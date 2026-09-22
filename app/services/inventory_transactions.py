"""Filtered inventory transaction listing and unlimited Excel export."""

from __future__ import annotations

from datetime import date, datetime, time, timedelta, timezone
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import func

from ..constants import EOD_TIMEZONE, LedgerType
from ..extensions import db
from ..models import (
    Carton,
    Client,
    InventoryTransaction,
    InventoryUnit,
    Order,
    PickTicket,
    Warehouse,
)
from .end_of_day import ny_day_utc_bounds
from .tenant import user_can_access_client


TZ = ZoneInfo(EOD_TIMEZONE)
PER_PAGE_OPTIONS = (100, 250, 500)
DEFAULT_PER_PAGE = 100
WORKSHEET_NAME = "Transactions"

EXPORT_COLUMNS = [
    "Timestamp",
    "Client",
    "Warehouse",
    "Unit ID",
    "UPC",
    "SKU",
    "Description",
    "Style",
    "Color",
    "Size",
    "Location",
    "Transaction",
    "From Status",
    "To Status",
    "WMS Order ID",
    "Client Order Number",
    "Customer",
    "Pick Ticket",
    "Carton Number",
    "Tracking Number",
    "Reference",
]


class TransactionAccessError(LookupError):
    """Client/warehouse scope is unauthorized."""


def parse_per_page(value) -> int:
    try:
        parsed = int(value)
    except (TypeError, ValueError):
        return DEFAULT_PER_PAGE
    return parsed if parsed in PER_PAGE_OPTIONS else DEFAULT_PER_PAGE


def parse_page(value) -> int:
    try:
        parsed = int(value)
    except (TypeError, ValueError):
        return 1
    return parsed if parsed > 0 else 1


def parse_day(value: str | None) -> date | None:
    value = (value or "").strip()
    if not value:
        return None
    return datetime.strptime(value, "%Y-%m-%d").date()


def to_ny(dt: datetime | None) -> datetime | None:
    if dt is None:
        return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(TZ)


def format_timestamp(dt: datetime | None) -> str:
    local = to_ny(dt)
    return local.strftime("%Y-%m-%d %H:%M:%S") if local else ""


def resolve_transaction_scope(user, client_id=None, warehouse_id=None):
    if client_id and not user_can_access_client(user, int(client_id)):
        raise TransactionAccessError("Client was not found.")
    warehouse = None
    if warehouse_id:
        warehouse = db.session.get(Warehouse, int(warehouse_id))
        if warehouse is None:
            warehouse_id = None
        elif client_id and warehouse.client_id != int(client_id):
            warehouse_id = None
        elif not user_can_access_client(user, warehouse.client_id):
            raise TransactionAccessError("Warehouse was not found.")
        else:
            warehouse_id = warehouse.id
            if not client_id:
                client_id = warehouse.client_id
    return client_id, warehouse_id


def ledger_query(
    *,
    client_id=None,
    warehouse_id=None,
    client_ids=None,
    upc=None,
    transaction_type=None,
    date_from=None,
    date_to=None,
):
    query = db.session.query(InventoryTransaction)
    if client_id:
        query = query.filter(InventoryTransaction.client_id == int(client_id))
    elif client_ids is not None:
        query = query.filter(InventoryTransaction.client_id.in_(client_ids or [-1]))
    if warehouse_id:
        query = query.filter(InventoryTransaction.warehouse_id == int(warehouse_id))
    if upc:
        query = query.filter(InventoryTransaction.upc == upc.strip())
    if transaction_type:
        kind = transaction_type.strip().upper()
        if kind in LedgerType.ALL:
            query = query.filter(InventoryTransaction.transaction_type == kind)
    if date_from is not None:
        query = query.filter(InventoryTransaction.created_at >= date_from)
    if date_to is not None:
        query = query.filter(InventoryTransaction.created_at < date_to)
    return query.order_by(InventoryTransaction.created_at.desc(), InventoryTransaction.id.desc())


def apply_screen_filters(user, filters: dict | None = None):
    filters = filters or {}
    client_id = filters.get("client_id") or None
    warehouse_id = filters.get("warehouse_id") or None
    client_id, warehouse_id = resolve_transaction_scope(user, client_id, warehouse_id)
    admin = bool(user and user.is_admin())
    accessible = list(user.client_ids() or []) if user else []
    client_ids = None if admin and not client_id else ([int(client_id)] if client_id else accessible)
    day_from = parse_day(filters.get("date_from"))
    day_to = parse_day(filters.get("date_to"))
    start = ny_day_utc_bounds(day_from)[0] if day_from else None
    end = ny_day_utc_bounds(day_to)[1] if day_to else None
    query = ledger_query(
        client_id=client_id,
        warehouse_id=warehouse_id,
        client_ids=client_ids,
        upc=filters.get("upc") or None,
        transaction_type=filters.get("type") or None,
        date_from=start,
        date_to=end,
    )
    return query, {
        "client_id": client_id,
        "warehouse_id": warehouse_id,
        "client_ids": client_ids,
        "upc": (filters.get("upc") or "").strip(),
        "type": (filters.get("type") or "").strip().upper(),
        "date_from": (filters.get("date_from") or "").strip(),
        "date_to": (filters.get("date_to") or "").strip(),
    }


def _maps_for(transactions: list[InventoryTransaction]) -> dict:
    client_ids = {row.client_id for row in transactions if row.client_id}
    warehouse_ids = {row.warehouse_id for row in transactions if row.warehouse_id}
    unit_ids = {row.inventory_unit_id for row in transactions if row.inventory_unit_id}
    order_ids = {row.order_id for row in transactions if row.order_id}
    ticket_ids = {row.pick_ticket_id for row in transactions if row.pick_ticket_id}
    carton_ids = {row.carton_id for row in transactions if row.carton_id}
    clients = {
        row.id: row
        for row in Client.query.filter(Client.id.in_(client_ids or [-1])).all()
    } if client_ids else {}
    warehouses = {
        row.id: row
        for row in Warehouse.query.filter(Warehouse.id.in_(warehouse_ids or [-1])).all()
    } if warehouse_ids else {}
    units = {
        row.id: row
        for row in InventoryUnit.query.filter(InventoryUnit.id.in_(unit_ids or [-1])).all()
    } if unit_ids else {}
    orders = {
        row.id: row
        for row in Order.query.filter(Order.id.in_(order_ids or [-1])).all()
    } if order_ids else {}
    tickets = {
        row.id: row
        for row in PickTicket.query.filter(PickTicket.id.in_(ticket_ids or [-1])).all()
    } if ticket_ids else {}
    cartons = {
        row.id: row
        for row in Carton.query.filter(Carton.id.in_(carton_ids or [-1])).all()
    } if carton_ids else {}
    return {
        "clients": clients,
        "warehouses": warehouses,
        "units": units,
        "orders": orders,
        "tickets": tickets,
        "cartons": cartons,
    }


def enrich_transactions(transactions: list[InventoryTransaction]) -> list[dict]:
    maps = _maps_for(transactions)
    rows = []
    for txn in transactions:
        unit = maps["units"].get(txn.inventory_unit_id)
        order = maps["orders"].get(txn.order_id)
        carton = maps["cartons"].get(txn.carton_id)
        ticket = maps["tickets"].get(txn.pick_ticket_id)
        client = maps["clients"].get(txn.client_id)
        warehouse = maps["warehouses"].get(txn.warehouse_id)
        rows.append(
            {
                "id": txn.id,
                "timestamp": format_timestamp(txn.created_at),
                "created_at": txn.created_at,
                "unit_id": txn.inventory_unit_id,
                "upc": txn.upc,
                "client": client.client_code if client else "",
                "warehouse": warehouse.warehouse_code if warehouse else "",
                "location": txn.location or "",
                "transaction": txn.transaction_type,
                "from_status": txn.from_status or "",
                "to_status": txn.to_status or "",
                "reference": txn.reference or "",
                "sku": unit.sku if unit else "",
                "description": unit.description if unit else "",
                "style": unit.style if unit else "",
                "color": unit.color if unit else "",
                "size": unit.size if unit else "",
                "wms_order_id": order.wms_order_id if order else "",
                "client_order_number": order.client_order_number if order else "",
                "customer": order.customer if order else "",
                "pick_ticket": ticket.pick_ticket_number if ticket else "",
                "carton_number": carton.carton_number if carton else "",
                "tracking_number": (carton.tracking_number or "") if carton else "",
            }
        )
    return rows


def list_transaction_page(user, filters: dict | None = None, *, page=1, per_page=DEFAULT_PER_PAGE):
    page = parse_page(page)
    per_page = parse_per_page(per_page)
    query, resolved = apply_screen_filters(user, filters)
    total = query.order_by(None).with_entities(func.count(InventoryTransaction.id)).scalar() or 0
    pages = max(1, (int(total) + per_page - 1) // per_page)
    if page > pages:
        page = pages
    transactions = query.offset((page - 1) * per_page).limit(per_page).all()
    return {
        "rows": enrich_transactions(transactions),
        "page": page,
        "per_page": per_page,
        "pages": pages,
        "total": int(total),
        "per_page_options": PER_PAGE_OPTIONS,
        "filters": resolved,
    }


def export_row_values(row: dict) -> list:
    return [
        row["timestamp"],
        row["client"],
        row["warehouse"],
        row["unit_id"],
        row["upc"],
        row["sku"] or "",
        row["description"] or "",
        row["style"] or "",
        row["color"] or "",
        row["size"] or "",
        row["location"],
        row["transaction"],
        row["from_status"],
        row["to_status"],
        row["wms_order_id"] or "",
        row["client_order_number"] or "",
        row["customer"] or "",
        row["pick_ticket"] or "",
        row["carton_number"] or "",
        row["tracking_number"] or "",
        row["reference"] or "",
    ]


def export_filename(when: datetime | None = None) -> str:
    stamp = to_ny(when or datetime.now(timezone.utc)) or datetime.now(TZ)
    return f"Inventory_Transactions_{stamp.strftime('%Y%m%d_%H%M%S')}.xlsx"


def export_transactions_xlsx(user, filters: dict | None = None) -> tuple[bytes, str, int]:
    query, _resolved = apply_screen_filters(user, filters)
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(WORKSHEET_NAME)
    header_font = Font(name="Arial", bold=True, color="182230")
    header_fill = PatternFill("solid", fgColor="F4F6F8")
    header_cells = []
    for title in EXPORT_COLUMNS:
        cell = WriteOnlyCell(ws, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left")
        header_cells.append(cell)
    ws.append(header_cells)
    exported = 0
    chunk: list[InventoryTransaction] = []
    for txn in query.yield_per(1000):
        chunk.append(txn)
        if len(chunk) >= 1000:
            for row in enrich_transactions(chunk):
                ws.append(export_row_values(row))
                exported += 1
            chunk.clear()
            db.session.expire_all()
    if chunk:
        for row in enrich_transactions(chunk):
            ws.append(export_row_values(row))
            exported += 1
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue(), export_filename(), exported


def export_transactions_csv_rows(user, filters: dict | None = None):
    query, _resolved = apply_screen_filters(user, filters)
    chunk: list[InventoryTransaction] = []
    for txn in query.yield_per(1000):
        chunk.append(txn)
        if len(chunk) >= 1000:
            yield from enrich_transactions(chunk)
            chunk.clear()
            db.session.expire_all()
    if chunk:
        yield from enrich_transactions(chunk)
