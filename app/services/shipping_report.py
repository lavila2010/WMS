"""Single-sheet Shipping Report: operational order list and Excel export."""

from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime, time, timedelta, timezone
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import func, or_
from sqlalchemy.orm import joinedload

from ..auth import record_audit
from ..constants import EOD_TIMEZONE, OrderStatus, ShippingStatus, TrackingCarrier
from ..extensions import db
from ..models import Carton, Order, PickTicket
from .carton_manifest import (
    FULFILLMENT_CLOSED_COMPLETE,
    FULFILLMENT_CLOSED_SHORT,
    FULFILLMENT_PARTIALLY,
    FULFILLMENT_STATUSES,
    display_status,
    fulfillment_status,
    line_totals,
    load_carton_manifests,
)
from .order_visibility import apply_operational_order_visibility
from .tenant import user_can_access_client


TZ = ZoneInfo(EOD_TIMEZONE)
PER_PAGE_OPTIONS = (25, 50, 100)
DEFAULT_PER_PAGE = 50
WORKSHEET_NAME = "Shipping Report"
REPORT_WIDTH = 17

# Fulfillment labels live in carton_manifest and are re-exported for callers.

ORDER_INFO_HEADERS = [
    "Client",
    "Order #",
    "Customer Name",
    "WMS Order ID",
    "Division",
    "Warehouse",
    "Carrier",
    "Shipping Service",
    "Order Status",
    "Fulfillment Status",
    "Shipping Status",
    "Ordered",
    "Shipped",
    "Short",
    "Cartons",
    "Created Date",
    "Closed Date",
]

PRODUCT_HEADERS = ["UPC", "SKU", "Description", "Style", "Color", "Size", "Qty"]

FONT_NAME = "Arial"
THIN = Border(
    left=Side(style="thin", color="E3E7EB"),
    right=Side(style="thin", color="E3E7EB"),
    top=Side(style="thin", color="E3E7EB"),
    bottom=Side(style="thin", color="E3E7EB"),
)
FILL_WHITE = PatternFill("solid", fgColor="FFFFFF")
FILL_HEADER = PatternFill("solid", fgColor="F4F6F8")
FILL_LABEL = PatternFill("solid", fgColor="EEF5FA")
FILL_CARTON = PatternFill("solid", fgColor="EAF7F0")
FONT_BASE = Font(name=FONT_NAME, size=10, color="182230")
FONT_BOLD = Font(name=FONT_NAME, size=10, bold=True, color="182230")
FONT_LABEL = Font(name=FONT_NAME, size=12, bold=True, color="182230")
ALIGN_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)


class ShippingReportAccessError(LookupError):
    """Selected order is missing or not authorized. Fail closed."""


class ShippingReportSelectionError(ValueError):
    """No orders were selected for export."""


def parse_per_page(value) -> int:
    try:
        parsed = int(value)
    except (TypeError, ValueError):
        return DEFAULT_PER_PAGE
    return parsed if parsed in PER_PAGE_OPTIONS else DEFAULT_PER_PAGE


def parse_page(value) -> int:
    try:
        parsed = int(value)
    except (TypeError, ValueError):
        return 1
    return parsed if parsed > 0 else 1


def parse_report_date(value: str | None) -> date | None:
    value = (value or "").strip()
    if not value:
        return None
    return datetime.strptime(value, "%Y-%m-%d").date()


def ny_day_utc_bounds(day: date) -> tuple[datetime, datetime]:
    start_local = datetime.combine(day, time.min, tzinfo=TZ)
    end_local = start_local + timedelta(days=1)
    start_utc = start_local.astimezone(timezone.utc).replace(tzinfo=None)
    end_utc = end_local.astimezone(timezone.utc).replace(tzinfo=None)
    return start_utc, end_utc


def to_ny(dt: datetime | None) -> datetime | None:
    if dt is None:
        return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(TZ)


def format_report_date(dt: datetime | None) -> str:
    local = to_ny(dt)
    return local.strftime("%Y-%m-%d") if local else ""


def order_label(initials: str | None, client_order_number: str | None, customer: str | None) -> str:
    parts = [str(initials or "").strip(), str(client_order_number or "").strip()]
    name = str(customer or "").strip()
    if name:
        parts.append(name)
    return " - ".join(part for part in parts if part)


def _unique_order_ids(raw_ids) -> list[int]:
    wanted: list[int] = []
    seen: set[int] = set()
    for raw in raw_ids or []:
        try:
            oid = int(raw)
        except (TypeError, ValueError):
            raise ShippingReportAccessError("Order was not found.")
        if oid in seen:
            continue
        seen.add(oid)
        wanted.append(oid)
    return wanted


def _apply_filters(query, filters: dict, *, accessible_client_ids, admin: bool):
    client_id = filters.get("client_id")
    division_id = filters.get("division_id")
    warehouse_id = filters.get("warehouse_id")
    order_status = (filters.get("order_status") or "").strip()
    fulfillment = (filters.get("fulfillment_status") or "").strip()
    shipping_status = (filters.get("shipping_status") or "").strip()
    carrier = (filters.get("carrier") or "").strip()
    q = (filters.get("q") or "").strip()

    if client_id:
        query = query.filter(Order.client_id == int(client_id))
    elif not admin:
        query = query.filter(Order.client_id.in_(accessible_client_ids or [-1]))
    if division_id:
        query = query.filter(Order.division_id == int(division_id))
    if warehouse_id:
        query = query.filter(Order.warehouse_id == int(warehouse_id))
    if order_status in OrderStatus.ALL:
        query = query.filter(Order.status == order_status)
    if fulfillment == FULFILLMENT_PARTIALLY:
        query = query.filter(Order.status == OrderStatus.PARTIALLY_FULFILLED)
    elif fulfillment == FULFILLMENT_CLOSED_COMPLETE:
        query = query.filter(Order.status == OrderStatus.CLOSED, Order.short_closed.is_(False))
    elif fulfillment == FULFILLMENT_CLOSED_SHORT:
        query = query.filter(Order.status == OrderStatus.CLOSED, Order.short_closed.is_(True))
    if shipping_status in ShippingStatus.ALL:
        query = query.filter(Order.shipping_status == shipping_status)
    if carrier:
        query = query.filter(Order.carrier.ilike(carrier))

    created_from = parse_report_date(filters.get("created_from"))
    created_to = parse_report_date(filters.get("created_to"))
    closed_from = parse_report_date(filters.get("closed_from"))
    closed_to = parse_report_date(filters.get("closed_to"))
    if created_from:
        start, _ = ny_day_utc_bounds(created_from)
        query = query.filter(Order.created_at >= start)
    if created_to:
        _, end = ny_day_utc_bounds(created_to)
        query = query.filter(Order.created_at < end)
    if closed_from:
        start, _ = ny_day_utc_bounds(closed_from)
        query = query.filter(Order.closed_at >= start)
    if closed_to:
        _, end = ny_day_utc_bounds(closed_to)
        query = query.filter(Order.closed_at < end)

    if q:
        like = f"%{q}%"
        ticket_match = (
            db.session.query(PickTicket.id)
            .filter(PickTicket.order_id == Order.id, PickTicket.pick_ticket_number.ilike(like))
            .exists()
        )
        tracking_match = (
            db.session.query(Carton.id)
            .filter(Carton.order_id == Order.id, Carton.tracking_number.ilike(like))
            .exists()
        )
        query = query.filter(
            or_(
                Order.wms_order_id.ilike(like),
                Order.client_order_number.ilike(like),
                Order.customer.ilike(like),
                ticket_match,
                tracking_match,
            )
        )
    return query


def shipping_report_query(user, filters: dict | None = None):
    filters = filters or {}
    admin = bool(user and user.is_admin())
    accessible = list(user.client_ids() or []) if user else []
    if filters.get("client_id") and not user_can_access_client(user, int(filters["client_id"])):
        raise ShippingReportAccessError("Order was not found.")
    query = apply_operational_order_visibility(Order.query)
    return _apply_filters(query, filters, accessible_client_ids=accessible, admin=admin)


def _line_totals(order_ids: list[int]) -> dict[int, dict[str, int]]:
    return line_totals(order_ids)


def _carton_summaries(order_ids: list[int]) -> dict[int, dict]:
    summaries: dict[int, dict] = defaultdict(lambda: {"carton_count": 0, "trackings": []})
    if not order_ids:
        return summaries
    rows = (
        db.session.query(Carton.order_id, Carton.tracking_number)
        .filter(Carton.order_id.in_(order_ids))
        .order_by(Carton.order_id.asc(), Carton.id.asc())
        .all()
    )
    for order_id, tracking in rows:
        bucket = summaries[int(order_id)]
        bucket["carton_count"] += 1
        number = (tracking or "").strip()
        if number and number not in bucket["trackings"]:
            bucket["trackings"].append(number)
    return summaries


def list_shipping_report_page(user, filters: dict | None = None, *, page=1, per_page=DEFAULT_PER_PAGE):
    page = parse_page(page)
    per_page = parse_per_page(per_page)
    query = shipping_report_query(user, filters)
    total = query.with_entities(func.count(Order.id)).scalar() or 0
    pages = max(1, (int(total) + per_page - 1) // per_page)
    if page > pages:
        page = pages
    orders = (
        query.options(
            joinedload(Order.client),
            joinedload(Order.division),
            joinedload(Order.warehouse),
        )
        .order_by(Order.created_at.desc(), Order.id.desc())
        .offset((page - 1) * per_page)
        .limit(per_page)
        .all()
    )
    ids = [order.id for order in orders]
    lines = _line_totals(ids)
    cartons = _carton_summaries(ids)
    rows = []
    for order in orders:
        qty = lines.get(order.id, {"ordered": 0, "shipped": 0, "short": 0})
        carton = cartons.get(order.id, {"carton_count": 0, "trackings": []})
        initials = order.client.initials if order.client else ""
        rows.append(
            {
                "order_id": order.id,
                "client": initials,
                "client_code": order.client.client_code if order.client else "",
                "order_number": order.client_order_number,
                "customer": order.customer or "",
                "order_label": order_label(initials, order.client_order_number, order.customer),
                "division": order.division.name if order.division else "",
                "warehouse": order.warehouse.warehouse_code if order.warehouse else "",
                "order_status": order.status,
                "fulfillment_status": fulfillment_status(order),
                "shipping_status": order.shipping_status,
                "ordered": qty["ordered"],
                "shipped": qty["shipped"],
                "short": qty["short"],
                "cartons": carton["carton_count"],
                "carrier": order.carrier or "",
                "tracking_summary": ", ".join(carton["trackings"]),
                "created_date": format_report_date(order.created_at),
                "closed_date": format_report_date(order.closed_at),
            }
        )
    return {
        "rows": rows,
        "page": page,
        "per_page": per_page,
        "total": int(total),
        "pages": pages,
        "per_page_options": PER_PAGE_OPTIONS,
    }


def _authorize_selected_orders(user, order_ids) -> list[Order]:
    wanted = _unique_order_ids(order_ids)
    if not wanted:
        raise ShippingReportSelectionError("Select at least one order to export.")
    allowed = set(user.client_ids() or [])
    found = {
        order.id: order
        for order in apply_operational_order_visibility(
            Order.query.filter(Order.id.in_(wanted)).options(
                joinedload(Order.client),
                joinedload(Order.division),
                joinedload(Order.warehouse),
            )
        ).all()
    }
    orders = []
    for oid in wanted:
        order = found.get(oid)
        if order is None or order.client_id not in allowed:
            raise ShippingReportAccessError("Order was not found.")
        orders.append(order)
    return orders


def load_shipping_report_blocks(user, order_ids) -> list[dict]:
    orders = _authorize_selected_orders(user, order_ids)
    order_ids = [order.id for order in orders]
    totals = _line_totals(order_ids)
    manifests = load_carton_manifests(order_ids)

    blocks = []
    for order in orders:
        qty = totals.get(order.id, {"ordered": 0, "shipped": 0, "short": 0})
        carton_blocks = manifests.get(order.id, [])
        initials = order.client.initials if order.client else ""
        blocks.append(
            {
                "order": order,
                "label": order_label(initials, order.client_order_number, order.customer),
                "client": initials,
                "order_number": order.client_order_number or "",
                "customer": order.customer or "",
                "wms_order_id": order.wms_order_id or "",
                "division": order.division.name if order.division else "",
                "warehouse": order.warehouse.warehouse_code if order.warehouse else "",
                "carrier": order.carrier or "",
                "shipping_service": order.shipping_service or "",
                "order_status": display_status(order.status),
                "fulfillment_status": fulfillment_status(order),
                "shipping_status": display_status(order.shipping_status),
                "ordered": qty["ordered"],
                "shipped": qty["shipped"],
                "short": qty["short"],
                "carton_count": len(carton_blocks),
                "created_date": format_report_date(order.created_at),
                "closed_date": format_report_date(order.closed_at),
                "cartons": carton_blocks,
            }
        )
    return blocks


def _style_range(ws, row: int, start=1, end=REPORT_WIDTH, *, font=None, fill=None, border=True):
    for col in range(start, end + 1):
        cell = ws.cell(row, col)
        cell.font = font or FONT_BASE
        cell.fill = fill or FILL_WHITE
        cell.alignment = ALIGN_WRAP
        if border:
            cell.border = THIN


def _write_merged(ws, row: int, value: str, fill, font):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=REPORT_WIDTH)
    cell = ws.cell(row, 1, value)
    cell.font = font
    cell.fill = fill
    cell.alignment = ALIGN_WRAP
    _style_range(ws, row, fill=fill, font=font)
    ws.cell(row, 1).value = value


def build_shipping_report_workbook(blocks: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = WORKSHEET_NAME
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    row = 1
    for block in blocks:
        _write_merged(ws, row, block["label"], FILL_LABEL, FONT_LABEL)
        ws.row_dimensions[row].height = 20
        row += 1
        for col, header in enumerate(ORDER_INFO_HEADERS, 1):
            cell = ws.cell(row, col, header)
            cell.font = FONT_BOLD
            cell.fill = FILL_HEADER
            cell.alignment = ALIGN_WRAP
            cell.border = THIN
        _style_range(ws, row, end=len(ORDER_INFO_HEADERS), font=FONT_BOLD, fill=FILL_HEADER)
        row += 1
        values = [
            block["client"],
            block["order_number"],
            block["customer"],
            block["wms_order_id"],
            block["division"],
            block["warehouse"],
            block["carrier"],
            block["shipping_service"],
            block["order_status"],
            block["fulfillment_status"],
            block["shipping_status"],
            block["ordered"],
            block["shipped"],
            block["short"],
            block["carton_count"],
            block["created_date"],
            block["closed_date"],
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row, col, value)
            if col in {12, 13, 14, 15} and value != "":
                cell.number_format = "0"
            cell.font = FONT_BASE
            cell.fill = FILL_WHITE
            cell.alignment = ALIGN_WRAP
            cell.border = THIN
        row += 1
        row += 1
        if not block["cartons"]:
            ws.cell(row, 1, "No cartons available.")
            _style_range(ws, row, font=FONT_BASE, fill=FILL_WHITE, border=False)
            row += 1
            row += 1
            row += 1
            continue
        for carton in block["cartons"]:
            _write_merged(ws, row, carton["header"], FILL_CARTON, FONT_BOLD)
            row += 1
            if not carton["products"]:
                ws.cell(row, 1, "No products")
                _style_range(ws, row, font=FONT_BASE, fill=FILL_WHITE, border=False)
                row += 1
            else:
                for col, header in enumerate(PRODUCT_HEADERS, 1):
                    cell = ws.cell(row, col, header)
                    cell.font = FONT_BOLD
                    cell.fill = FILL_HEADER
                    cell.alignment = ALIGN_WRAP
                    cell.border = THIN
                row += 1
                for product in carton["products"]:
                    product_values = [
                        product["upc"],
                        product["sku"],
                        product["description"],
                        product["style"],
                        product["color"],
                        product["size"],
                        product["qty"],
                    ]
                    for col, value in enumerate(product_values, 1):
                        cell = ws.cell(row, col, value)
                        if col == 7:
                            cell.number_format = "0"
                        cell.font = FONT_BASE
                        cell.fill = FILL_WHITE
                        cell.alignment = ALIGN_WRAP
                        cell.border = THIN
                    row += 1
            row += 1
        row += 1

    widths = [18, 14, 22, 20, 14, 16, 12, 16, 16, 18, 20, 10, 10, 10, 10, 14, 14]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def export_filename(when: datetime | None = None) -> str:
    stamp = to_ny(when or datetime.now(timezone.utc)) or datetime.now(TZ)
    return f"Shipping_Report_{stamp.strftime('%Y%m%d_%H%M%S')}.xlsx"


def export_shipping_report(user, order_ids) -> tuple[bytes, str, dict]:
    blocks = load_shipping_report_blocks(user, order_ids)
    product_rows = sum(len(carton["products"]) for block in blocks for carton in block["cartons"])
    carton_count = sum(len(block["cartons"]) for block in blocks)
    data = build_shipping_report_workbook(blocks)
    client_ids = {block["order"].client_id for block in blocks}
    record_audit(
        "SHIPPING_REPORT_EXPORTED",
        module="Orders",
        entity_type="shipping_report",
        entity_id=blocks[0]["order"].id if len(blocks) == 1 else None,
        client_id=next(iter(client_ids)) if len(client_ids) == 1 else None,
        detail=f"selected_orders={len(blocks)} cartons={carton_count} product_rows={product_rows}",
    )
    stats = {
        "selected_orders": len(blocks),
        "cartons": carton_count,
        "product_rows": product_rows,
    }
    return data, export_filename(), stats


def report_filter_choices():
    return {
        "order_statuses": OrderStatus.ALL,
        "fulfillment_statuses": FULFILLMENT_STATUSES,
        "shipping_statuses": ShippingStatus.ALL,
        "carriers": TrackingCarrier.ALL,
        "per_page_options": PER_PAGE_OPTIONS,
    }
