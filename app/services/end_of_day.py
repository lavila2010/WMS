"""End of Day fulfillment/shipping manifest in America/New_York operational time."""

from __future__ import annotations

from datetime import date, datetime, time, timedelta, timezone
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import and_, exists, or_
from sqlalchemy.orm import joinedload

from ..auth import record_audit
from ..constants import CartonStatus, EOD_TIMEZONE, OrderStatus, ShippingStatus
from ..extensions import db
from ..models import Carton, Client, Division, Order
from .carton_manifest import (
    FULFILLMENT_CLOSED_COMPLETE,
    FULFILLMENT_CLOSED_SHORT,
    FULFILLMENT_PARTIALLY,
    FULFILLMENT_STATUSES,
    carton_in_day,
    fulfillment_status,
    is_processed_carton,
    line_totals,
    load_carton_manifests,
    load_cartons_by_order,
    remaining_units,
)
from .order_visibility import apply_operational_order_visibility
from .shipping_report import PRODUCT_HEADERS, order_label
from .tenant import user_can_access_client


TZ = ZoneInfo(EOD_TIMEZONE)
WORKSHEET_NAME = "End of Day"
REPORT_WIDTH = 20
PRODUCT_HEADERS = PRODUCT_HEADERS

ORDER_INFO_HEADERS = [
    "Client",
    "Order #",
    "Customer Name",
    "WMS Order ID",
    "Division",
    "Warehouse",
    "Carrier",
    "Shipping Service",
    "Order Status",
    "Fulfillment Status",
    "Shipping Status",
    "Ordered",
    "Shipped",
    "Short",
    "Remaining",
    "Cartons Processed Today",
    "Total Cartons",
    "Total Weight",
    "Processed / Closed Date",
    "Processed / Closed Time",
]

FONT_NAME = "Arial"
THIN = Border(
    left=Side(style="thin", color="E3E7EB"),
    right=Side(style="thin", color="E3E7EB"),
    top=Side(style="thin", color="E3E7EB"),
    bottom=Side(style="thin", color="E3E7EB"),
)
FILL_WHITE = PatternFill("solid", fgColor="FFFFFF")
FILL_HEADER = PatternFill("solid", fgColor="F4F6F8")
FILL_LABEL = PatternFill("solid", fgColor="EEF5FA")
FILL_CARTON = PatternFill("solid", fgColor="EAF7F0")
FONT_BASE = Font(name=FONT_NAME, size=10, color="182230")
FONT_BOLD = Font(name=FONT_NAME, size=10, bold=True, color="182230")
FONT_LABEL = Font(name=FONT_NAME, size=12, bold=True, color="182230")
ALIGN_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)


def operational_today() -> date:
    return datetime.now(TZ).date()


def ny_day_utc_bounds(day: date) -> tuple[datetime, datetime]:
    start_local = datetime.combine(day, time.min, tzinfo=TZ)
    end_local = start_local + timedelta(days=1)
    start_utc = start_local.astimezone(timezone.utc).replace(tzinfo=None)
    end_utc = end_local.astimezone(timezone.utc).replace(tzinfo=None)
    return start_utc, end_utc


def parse_eod_date(value: str | None) -> date:
    if not value:
        return operational_today()
    return datetime.strptime(value, "%Y-%m-%d").date()


def to_ny(dt: datetime | None) -> datetime | None:
    if dt is None:
        return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(TZ)


def closed_orders_query(
    *,
    day: date,
    client_id=None,
    division_id=None,
    warehouse_id=None,
    carrier="",
    closed_by="",
    shipping_status="",
    fulfillment_status_filter="",
    order_status="",
    accessible_client_ids=None,
    admin=False,
    user=None,
):
    start, end = ny_day_utc_bounds(day)
    today_carton = exists().where(
        Carton.order_id == Order.id,
        Carton.status == CartonStatus.CLOSED,
        Carton.closed_at >= start,
        Carton.closed_at < end,
    )
    closed = and_(
        Order.status == OrderStatus.CLOSED,
        Order.closed_at >= start,
        Order.closed_at < end,
    )
    partial = and_(Order.status == OrderStatus.PARTIALLY_FULFILLED, today_carton)
    fulfillment = (fulfillment_status_filter or "").strip()
    if fulfillment == FULFILLMENT_CLOSED_COMPLETE:
        predicate = and_(closed, Order.short_closed.is_(False))
    elif fulfillment == FULFILLMENT_CLOSED_SHORT:
        predicate = and_(closed, Order.short_closed.is_(True))
    elif fulfillment == FULFILLMENT_PARTIALLY:
        predicate = partial
    else:
        predicate = or_(closed, partial)
    query = apply_operational_order_visibility(Order.query.filter(predicate))
    if client_id:
        if user is not None and not user_can_access_client(user, int(client_id)):
            query = query.filter(Order.id == -1)
        else:
            query = query.filter(Order.client_id == int(client_id))
    elif not admin:
        query = query.filter(Order.client_id.in_(accessible_client_ids or [-1]))
    if division_id:
        query = query.filter(Order.division_id == int(division_id))
    if warehouse_id:
        query = query.filter(Order.warehouse_id == int(warehouse_id))
    if carrier:
        query = query.filter(Order.carrier.ilike(f"%{carrier.strip()}%"))
    if closed_by:
        query = query.filter(Order.closed_by_username.ilike(f"%{closed_by.strip()}%"))
    if shipping_status in (ShippingStatus.PENDING_TRACKING, ShippingStatus.TRACKING_COMPLETE, ShippingStatus.NOT_READY):
        query = query.filter(Order.shipping_status == shipping_status)
    if order_status in OrderStatus.ALL:
        query = query.filter(Order.status == order_status)
    return (
        query.join(Client, Client.id == Order.client_id)
        .join(Division, Division.id == Order.division_id)
        .options(
            joinedload(Order.client),
            joinedload(Order.division),
            joinedload(Order.warehouse),
        )
        .order_by(
            Client.client_code.asc(),
            Division.code.asc(),
            Order.closed_at.asc(),
            Order.id.asc(),
        )
    )


def _processed_at(order: Order, today_cartons: list, day: date | None):
    if order.status == OrderStatus.CLOSED and order.closed_at:
        return order.closed_at
    stamps = [carton.closed_at for carton in today_cartons if carton.closed_at]
    return max(stamps) if stamps else None


def _reconciled(status: str, ordered: int, shipped: int, short: int, remaining: int, carton_units: int) -> bool:
    if status == FULFILLMENT_CLOSED_COMPLETE:
        return shipped + short == ordered and short == 0 and remaining == 0
    if status == FULFILLMENT_CLOSED_SHORT:
        return shipped + short == ordered and short > 0
    if status == FULFILLMENT_PARTIALLY:
        return shipped + short < ordered and remaining > 0
    return True


def build_eod_rows(orders: list[Order], day: date | None = None) -> list[dict]:
    order_ids = [order.id for order in orders]
    totals = line_totals(order_ids)
    cartons_by_order = load_cartons_by_order(order_ids)
    start = end = None
    if day is not None:
        start, end = ny_day_utc_bounds(day)
    display_ids = []
    today_ids_by_order: dict[int, list[int]] = {}
    processed_by_order: dict[int, list] = {}
    for order in orders:
        all_cartons = cartons_by_order.get(order.id, [])
        processed = [carton for carton in all_cartons if is_processed_carton(carton)]
        processed_by_order[order.id] = processed
        if start is not None:
            today = [carton for carton in processed if carton_in_day(carton, start, end)]
        else:
            today = list(processed)
        today_ids_by_order[order.id] = [carton.id for carton in today]
        if order.status == OrderStatus.PARTIALLY_FULFILLED:
            display_ids.extend(carton.id for carton in today)
        else:
            display_ids.extend(carton.id for carton in processed)
    manifests = load_carton_manifests(order_ids, carton_ids=display_ids)

    rows = []
    for order in orders:
        qty = totals.get(
            order.id,
            {"ordered": 0, "shipped": 0, "short": 0, "remaining": 0, "allocated": 0, "packed": 0},
        )
        remaining = qty.get("remaining", remaining_units(qty["ordered"], qty["shipped"], qty["short"]))
        processed = processed_by_order.get(order.id, [])
        today_ids = set(today_ids_by_order.get(order.id, []))
        today_cartons = [carton for carton in processed if carton.id in today_ids]
        blocks = manifests.get(order.id, [])
        display_cartons = [block["carton"] for block in blocks]
        displayed_weight = sum((carton.weight or 0) for carton in display_cartons)
        displayed_units = sum(int(block["unit_count"]) for block in blocks)
        status = fulfillment_status(order)
        processed_at = _processed_at(order, today_cartons, day)
        local = to_ny(processed_at)
        rows.append(
            {
                "order": order,
                "label": order_label(
                    order.client.initials if order.client else "",
                    order.client_order_number,
                    order.customer,
                ),
                "fulfillment_status": status,
                "close_type": status,
                "short_closed": bool(getattr(order, "short_closed", False)),
                "short_reason": order.short_close_reason if getattr(order, "short_closed", False) else "",
                "ordered": qty["ordered"],
                "shipped": qty["shipped"],
                "short": qty["short"],
                "remaining": remaining,
                "allocated": qty.get("allocated", 0),
                "packed": qty.get("packed", 0),
                "cartons": display_cartons,
                "carton_blocks": blocks,
                "carton_count": len(processed),
                "cartons_processed_today": len(today_cartons),
                "total_cartons": len(processed),
                "carton_units": displayed_units,
                "total_weight": displayed_weight,
                "shipping_status": order.shipping_status or ShippingStatus.NOT_READY,
                "closed_local": local,
                "processed_local": local,
                "ticket": blocks[0]["pick_ticket_number"] if blocks else "—",
                "reconciled": _reconciled(status, qty["ordered"], qty["shipped"], qty["short"], remaining, displayed_units),
            }
        )
    return rows


def eod_kpis(rows: list[dict]) -> dict:
    return {
        "total_orders": len(rows),
        "closed_complete": sum(1 for r in rows if r["fulfillment_status"] == FULFILLMENT_CLOSED_COMPLETE),
        "closed_short": sum(1 for r in rows if r["fulfillment_status"] == FULFILLMENT_CLOSED_SHORT),
        "partially_fulfilled": sum(1 for r in rows if r["fulfillment_status"] == FULFILLMENT_PARTIALLY),
        "total_units": sum(r["shipped"] for r in rows),
        "short_units": sum(r["short"] for r in rows),
        "total_cartons": sum(r["cartons_processed_today"] for r in rows),
        "total_weight": sum(r["total_weight"] for r in rows),
        "tracking_complete": sum(1 for r in rows if r["shipping_status"] == ShippingStatus.TRACKING_COMPLETE),
        "pending_tracking": sum(1 for r in rows if r["shipping_status"] == ShippingStatus.PENDING_TRACKING),
    }


def _style_range(ws, row: int, start=1, end=REPORT_WIDTH, *, font=None, fill=None, border=True):
    for col in range(start, end + 1):
        cell = ws.cell(row, col)
        cell.font = font or FONT_BASE
        cell.fill = fill or FILL_WHITE
        cell.alignment = ALIGN_WRAP
        if border:
            cell.border = THIN


def _write_merged(ws, row: int, value: str, fill, font):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=REPORT_WIDTH)
    cell = ws.cell(row, 1, value)
    cell.font = font
    cell.fill = fill
    cell.alignment = ALIGN_WRAP
    _style_range(ws, row, fill=fill, font=font)
    ws.cell(row, 1).value = value


def _build_eod_workbook(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = WORKSHEET_NAME
    ws.sheet_view.showGridLines = False
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    row = 1
    for item in rows:
        order = item["order"]
        local = item["processed_local"]
        _write_merged(ws, row, item["label"], FILL_LABEL, FONT_LABEL)
        ws.row_dimensions[row].height = 20
        row += 1
        for col, header in enumerate(ORDER_INFO_HEADERS, 1):
            cell = ws.cell(row, col, header)
            cell.font = FONT_BOLD
            cell.fill = FILL_HEADER
            cell.alignment = ALIGN_WRAP
            cell.border = THIN
        row += 1
        values = [
            order.client.initials if order.client else "",
            order.client_order_number or "",
            order.customer or "",
            order.wms_order_id or "",
            order.division.code if order.division else "",
            order.warehouse.warehouse_code if order.warehouse else "",
            order.carrier or "",
            order.shipping_service or "",
            order.status or "",
            item["fulfillment_status"],
            item["shipping_status"],
            item["ordered"],
            item["shipped"],
            item["short"],
            item["remaining"],
            item["cartons_processed_today"],
            item["total_cartons"],
            item["total_weight"],
            local.strftime("%Y-%m-%d") if local else "",
            local.strftime("%H:%M:%S") if local else "",
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row, col, value)
            if col in {12, 13, 14, 15, 16, 17} and value != "":
                cell.number_format = "0"
            cell.font = FONT_BASE
            cell.fill = FILL_WHITE
            cell.alignment = ALIGN_WRAP
            cell.border = THIN
        row += 2
        if not item["carton_blocks"]:
            ws.cell(row, 1, "No cartons processed.")
            _style_range(ws, row, font=FONT_BASE, fill=FILL_WHITE, border=False)
            row += 3
            continue
        for block in item["carton_blocks"]:
            _write_merged(ws, row, block["header"], FILL_CARTON, FONT_BOLD)
            row += 1
            if not block["products"]:
                ws.cell(row, 1, "No products")
                _style_range(ws, row, font=FONT_BASE, fill=FILL_WHITE, border=False)
                row += 1
            else:
                for col, header in enumerate(PRODUCT_HEADERS, 1):
                    cell = ws.cell(row, col, header)
                    cell.font = FONT_BOLD
                    cell.fill = FILL_HEADER
                    cell.alignment = ALIGN_WRAP
                    cell.border = THIN
                row += 1
                for product in block["products"]:
                    product_values = [
                        product["upc"],
                        product["sku"],
                        product["description"],
                        product["style"],
                        product["color"],
                        product["size"],
                        product["qty"],
                    ]
                    for col, value in enumerate(product_values, 1):
                        cell = ws.cell(row, col, value)
                        if col == 7:
                            cell.number_format = "0"
                        cell.font = FONT_BASE
                        cell.fill = FILL_WHITE
                        cell.alignment = ALIGN_WRAP
                        cell.border = THIN
                    row += 1
            row += 1
        row += 1
    widths = [16, 12, 20, 20, 12, 16, 12, 16, 20, 20, 18, 10, 10, 10, 12, 16, 12, 12, 16, 16]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def export_eod_excel(rows: list[dict], *, client_id=None, day: date | None = None) -> bytes:
    data = _build_eod_workbook(rows)
    kpis = eod_kpis(rows)
    record_audit(
        "END_OF_DAY_EXPORTED",
        module="Orders",
        entity_type="end_of_day",
        client_id=client_id,
        detail=(
            f"date={day.isoformat() if day else ''} orders={kpis['total_orders']} "
            f"closed_complete={kpis['closed_complete']} closed_short={kpis['closed_short']} "
            f"partially_fulfilled={kpis['partially_fulfilled']} cartons={kpis['total_cartons']} "
            f"shipped_units={kpis['total_units']}"
        ),
    )
    db.session.commit()
    return data
