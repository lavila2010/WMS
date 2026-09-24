"""Completed-orders Excel export for Order Management and Reports."""

from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import func

from ..constants import DocumentType, OrderStatus
from ..models import Document, Order
from .packing import packed_count


EXPORT_COLUMNS = [
    "Closed Date/Time",
    "Client",
    "Warehouse",
    "Order Type",
    "Order Number",
    "Pick Ticket Number",
    "Customer",
    "Carrier",
    "Shipping Service",
    "Ordered Units",
    "Packed Units",
    "Carton Count",
    "Total Weight",
    "Closed By",
    "Closure PDF",
]


def _parse_date(value):
    value = (value or "").strip()
    if not value:
        return None
    return datetime.strptime(value, "%Y-%m-%d").date()


def completed_orders_query(filters: dict | None = None):
    filters = filters or {}
    query = Order.query.filter(Order.status == OrderStatus.CLOSED)
    if filters.get("client_id"):
        query = query.filter(Order.client_id == int(filters["client_id"]))
    if filters.get("warehouse_id"):
        query = query.filter(Order.warehouse_id == int(filters["warehouse_id"]))
    if filters.get("order_type_id"):
        query = query.filter(Order.order_type_id == int(filters["order_type_id"]))
    carrier = (filters.get("carrier") or "").strip()
    if carrier:
        query = query.filter(Order.carrier.ilike(f"%{carrier}%"))
    closed_by = (filters.get("closed_by") or "").strip()
    if closed_by:
        query = query.filter(Order.closed_by_username.ilike(f"%{closed_by}%"))
    date_from = _parse_date(filters.get("date_from"))
    date_to = _parse_date(filters.get("date_to"))
    closed_expr = func.coalesce(Order.closed_at, Order.updated_at)
    if date_from:
        query = query.filter(func.date(closed_expr) >= date_from)
    if date_to:
        query = query.filter(func.date(closed_expr) <= date_to)
    return query.order_by(closed_expr.desc())


def _closure_pdf_name(order: Order) -> str:
    doc = (
        Document.query.filter_by(order_id=order.id, type=DocumentType.ORDER_CLOSURE)
        .order_by(Document.created_at.desc())
        .first()
    )
    return doc.filename if doc else ""


def order_export_row(order: Order) -> list:
    closed_at = order.closed_at or order.updated_at
    ticket = getattr(order, "pick_ticket", None)
    total_weight = round(sum(b.weight_kg or 0.0 for b in order.boxes), 3)
    unit = next((b.weight_unit for b in order.boxes if b.weight_unit), "lb")
    return [
        closed_at.strftime("%Y-%m-%d %H:%M") if closed_at else "",
        order.client.code if order.client else "",
        order.warehouse.code if order.warehouse else "",
        order.order_type.code if order.order_type else "",
        order.order_number,
        ticket.pick_ticket_number if ticket else "",
        order.customer or "",
        order.carrier or "",
        order.shipping_service or "",
        order.ordered_quantity,
        packed_count(order),
        len(order.boxes),
        f"{total_weight:g} {unit}".strip(),
        order.closed_by_username or "",
        _closure_pdf_name(order),
    ]


def build_completed_orders_workbook(orders) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Completed Orders"
    header_font = Font(bold=True, color="182230")
    header_fill = PatternFill("solid", fgColor="F4F6F8")
    for col, title in enumerate(EXPORT_COLUMNS, 1):
        cell = ws.cell(1, col, title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left")
    for row_idx, order in enumerate(orders, 2):
        for col, value in enumerate(order_export_row(order), 1):
            ws.cell(row_idx, col, value)
    for column in ws.columns:
        width = min(40, max(12, max(len(str(c.value or "")) for c in column) + 2))
        ws.column_dimensions[column[0].column_letter].width = width
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
