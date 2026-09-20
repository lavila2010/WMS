"""CELINE-scale inventory import: staging, bulk insert, visibility, idempotency."""

from __future__ import annotations

import io
import json
import time
from pathlib import Path

from openpyxl import Workbook

from app.constants import ImportBatchStatus, LedgerType, OrderStatus, UnitStatus
from app.models import ImportBatch, InventoryTransaction, InventoryUnit, Order
from app.services.allocation import allocate_order
from app.services.inventory_import import (
    analyze,
    cleanup_failed_import,
    commit_import,
    normalize_location,
    normalize_upc,
    preview_from_batch,
    process_import_batch,
    retry_import,
)
from app.services.inventory_query import status_counts
from app.services.invariants import assert_invariants
from tests.test_v2_phase02_inventory import _masters, _row, _xlsx
from tests.test_v2_phase03_orders import _line
from tests.test_v2_phase04_allocation import _import_order, _world

CELINE_ROWS = 18436
CELINE_UNITS = 50336
CELINE_UPCS = 6638
CELINE_STYLES = 2312
CELINE_LOCS = 2871
CAPACITY_ROWS = 25000
CAPACITY_UNITS = 150000
ARTIFACT = Path("/opt/cursor/artifacts/large_import_benchmark.json")


def _write_xlsx(rows) -> io.BytesIO:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet()
    sheet.append(["UPC", "SKU", "Description", "Style", "Color", "Size", "Quantity", "Location"])
    for row in rows:
        sheet.append(list(row))
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def _celine_rows():
    colors = ("BLK", "TAN", "RED", "NAV")
    sizes = ("XS", "S", "M", "L", "XL")
    rows = []
    qty_tail = [3] * 8361 + [2] * 10074
    assert len(qty_tail) == CELINE_ROWS - 1
    quantities = [5105] + qty_tail
    assert sum(quantities) == CELINE_UNITS
    for index, qty in enumerate(quantities):
        upc = f"{(index % CELINE_UPCS):012d}"
        style = f"ST-{(index % CELINE_STYLES):04d}"
        location = f"{(index % CELINE_LOCS):010d}/0001"
        rows.append(
            (
                upc,
                "",
                f"Celine Item {(index % CELINE_UPCS):04d}",
                style,
                colors[index % len(colors)],
                sizes[index % len(sizes)],
                qty,
                location,
            )
        )
    return rows


def _capacity_rows():
    rows = []
    for index in range(CAPACITY_ROWS):
        rows.append(
            (
                f"{(index % 8000):013d}",
                "",
                f"Capacity Item {index % 8000}",
                f"ST-{(index % 2000):04d}",
                "BLK",
                "M",
                6,
                f"{(index % 3000):010d}/0002",
            )
        )
    return rows


def _write_benchmark(payload: dict) -> None:
    ARTIFACT.parent.mkdir(parents=True, exist_ok=True)
    existing = {}
    if ARTIFACT.is_file():
        existing = json.loads(ARTIFACT.read_text())
    existing.update(payload)
    ARTIFACT.write_text(json.dumps(existing, indent=2))


def test_blank_sku_accepted(app, db, admin_user):
    celine, _, cel_ny, _, _ = _masters(db)
    preview = analyze(_xlsx([_row(qty=2, sku="")]), "blank-sku.xlsx", celine, cel_ny)
    assert preview["has_blocking"] is False
    batch = commit_import(preview, user=admin_user)
    units = InventoryUnit.query.filter_by(import_batch_id=batch.id).all()
    assert len(units) == 2
    assert all(u.sku is None for u in units)


def test_description_required_blocks_file(app, db, admin_user):
    celine, _, cel_ny, _, _ = _masters(db)
    preview = analyze(_xlsx([_row(qty=1, description="")]), "nodesc.xlsx", celine, cel_ny)
    assert preview["has_blocking"] is True
    assert any("Description is required" in e["message"] for e in preview["blocking"])


def test_upc_leading_zeros_and_numeric_excel(app, db, admin_user):
    celine, _, cel_ny, _, _ = _masters(db)
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["UPC", "SKU", "Description", "Style", "Color", "Size", "Quantity", "Location"])
    sheet.append(["000123456789", "", "Zero Twelve", "ST", "BLK", "M", 1, "A-01"])
    sheet.append(["012345678905", "", "Zero Thirteen", "ST", "BLK", "M", 1, "A-02"])
    sheet.append([1234567890123, "", "Thirteen Digit", "ST", "BLK", "M", 1, "A-03"])
    sheet.append([1.234567890123e12, "", "Scientific", "ST", "BLK", "M", 1, "A-04"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    preview = analyze(buffer, "upc.xlsx", celine, cel_ny)
    assert preview["has_blocking"] is False, preview["blocking"]
    commit_import(preview, user=admin_user)
    upcs = {u.upc for u in InventoryUnit.query.all()}
    assert "000123456789" in upcs
    assert "012345678905" in upcs
    assert "1234567890123" in upcs
    assert all(not upc.endswith(".0") for upc in upcs)
    assert normalize_upc(1234567890123.0) == "1234567890123"
    assert normalize_upc("000123456789") == "000123456789"


def test_location_leading_zeros_preserved(app, db, admin_user):
    celine, _, cel_ny, _, _ = _masters(db)
    preview = analyze(
        _xlsx([_row(qty=1, location="0080097422/0001")]),
        "loc.xlsx",
        celine,
        cel_ny,
    )
    commit_import(preview, user=admin_user)
    unit = InventoryUnit.query.one()
    assert unit.location == "0080097422/0001"
    assert normalize_location("0080097422/0001") == "0080097422/0001"


def test_quantity_5105_chunked(app, db, admin_user):
    celine, _, cel_ny, _, _ = _masters(db)
    preview = analyze(_xlsx([_row(qty=5105, sku="")]), "big.xlsx", celine, cel_ny)
    assert preview["units"] == 5105
    batch = commit_import(preview, user=admin_user, chunk_size=2000)
    assert batch.status == ImportBatchStatus.COMPLETED
    assert InventoryUnit.query.filter_by(import_batch_id=batch.id).count() == 5105
    assert (
        InventoryTransaction.query.filter_by(
            import_batch_id=batch.id, transaction_type=LedgerType.IMPORT
        ).count()
        == 5105
    )


def test_one_import_ledger_per_unit(app, db, admin_user):
    celine, _, cel_ny, _, _ = _masters(db)
    batch = commit_import(
        analyze(_xlsx([_row(qty=7)]), "led.xlsx", celine, cel_ny),
        user=admin_user,
    )
    units = InventoryUnit.query.filter_by(import_batch_id=batch.id).count()
    txns = InventoryTransaction.query.filter_by(
        import_batch_id=batch.id, transaction_type=LedgerType.IMPORT
    ).count()
    assert units == txns == 7


def test_incomplete_and_failed_batch_hidden_from_allocation(app, db, admin_user):
    w = _world(db)
    preview2 = analyze(_xlsx([_row(upc="HIDE-2", qty=8)]), "hide2.xlsx", w["celine"], w["cel_ny"])
    mid_batch = ImportBatch.query.get(preview2["batch_id"])
    mid_batch.status = ImportBatchStatus.PROCESSING
    db.session.commit()
    process_import_batch(preview2["batch_id"], chunk_size=3, max_chunks=1)
    mid = ImportBatch.query.get(preview2["batch_id"])
    assert mid.status == ImportBatchStatus.PROCESSING
    assert mid.units_created == 3
    assert InventoryUnit.query.filter_by(import_batch_id=mid.id).count() == 3
    assert status_counts(client_id=w["celine"].id, warehouse_id=w["cel_ny"].id, upc="HIDE-2")["available"] == 0
    order = _import_order(admin_user, w["celine"], w["cel_ecom"], [_line(upc="HIDE-2", qty=2)])
    result = allocate_order(order, user=admin_user)
    assert result["reserved"] == 0
    assert Order.query.get(order.id).status == OrderStatus.UNALLOCATED

    preview3 = analyze(_xlsx([_row(upc="HIDE-3", qty=6)]), "hide3.xlsx", w["celine"], w["cel_ny"])
    try:
        commit_import(preview3, user=admin_user, chunk_size=2, _fail_after=2)
    except Exception:
        pass
    failed = ImportBatch.query.get(preview3["batch_id"])
    assert failed.status == ImportBatchStatus.FAILED
    assert status_counts(client_id=w["celine"].id, warehouse_id=w["cel_ny"].id, upc="HIDE-3")["available"] == 0
    order3 = _import_order(admin_user, w["celine"], w["cel_ecom"], [_line(order="3003", upc="HIDE-3", qty=1)])
    result3 = allocate_order(order3, user=admin_user)
    assert result3["reserved"] == 0


def test_double_confirm_and_retry_do_not_duplicate(app, db, admin_user):
    celine, _, cel_ny, _, _ = _masters(db)
    preview = analyze(_xlsx([_row(upc="IDEM-1", qty=6)]), "idem.xlsx", celine, cel_ny)
    first = commit_import(preview, user=admin_user)
    second = commit_import(preview, user=admin_user)
    assert first.id == second.id
    assert InventoryUnit.query.filter_by(import_batch_id=first.id).count() == 6
    assert (
        InventoryTransaction.query.filter_by(
            import_batch_id=first.id, transaction_type=LedgerType.IMPORT
        ).count()
        == 6
    )

    preview_fail = analyze(_xlsx([_row(upc="IDEM-2", qty=6)]), "retry.xlsx", celine, cel_ny)
    try:
        commit_import(preview_fail, user=admin_user, chunk_size=2, _fail_after=2)
    except Exception:
        pass
    failed = ImportBatch.query.get(preview_fail["batch_id"])
    assert failed.status == ImportBatchStatus.FAILED
    assert failed.units_created == 2
    retried = retry_import(failed.id, user=admin_user, run="sync")
    assert retried.status == ImportBatchStatus.COMPLETED
    assert InventoryUnit.query.filter_by(import_batch_id=failed.id).count() == 6
    assert (
        InventoryTransaction.query.filter_by(
            import_batch_id=failed.id, transaction_type=LedgerType.IMPORT
        ).count()
        == 6
    )
    cleanup = analyze(_xlsx([_row(upc="IDEM-3", qty=4)]), "clean.xlsx", celine, cel_ny)
    try:
        commit_import(cleanup, user=admin_user, chunk_size=2, _fail_after=2)
    except Exception:
        pass
    dirty = ImportBatch.query.get(cleanup["batch_id"])
    cleanup_failed_import(dirty.id, user=admin_user)
    assert InventoryUnit.query.filter_by(import_batch_id=dirty.id).count() == 0
    assert InventoryTransaction.query.filter_by(import_batch_id=dirty.id).count() == 0


def test_client_and_warehouse_isolation(app, db, admin_user):
    celine, dior, cel_ny, cel_nj, dio_ny = _masters(db)
    commit_import(analyze(_xlsx([_row(upc="ISO-1", qty=3)]), "c.xlsx", celine, cel_ny), user=admin_user)
    commit_import(analyze(_xlsx([_row(upc="ISO-1", qty=4)]), "d.xlsx", dior, dio_ny), user=admin_user)
    commit_import(analyze(_xlsx([_row(upc="ISO-1", qty=2, location="NJ-1")]), "n.xlsx", celine, cel_nj), user=admin_user)
    assert status_counts(client_id=celine.id, warehouse_id=cel_ny.id, upc="ISO-1")["available"] == 3
    assert status_counts(client_id=dior.id, warehouse_id=dio_ny.id, upc="ISO-1")["available"] == 4
    assert status_counts(client_id=celine.id, warehouse_id=cel_nj.id, upc="ISO-1")["available"] == 2
    assert InventoryUnit.query.filter_by(client_id=celine.id, warehouse_id=dio_ny.id).count() == 0


def test_product_attributes_preserved(app, db, admin_user):
    celine, _, cel_ny, _, _ = _masters(db)
    row = _row(
        upc="000999000111",
        qty=3,
        sku="SKU-KEEP",
        description="Full Attribute Coat",
        style="STYLE-Z",
        color="BURGUNDY",
        size="XXL",
        location="0080097422/0001",
    )
    commit_import(analyze(_xlsx([row]), "attr.xlsx", celine, cel_ny), user=admin_user)
    units = InventoryUnit.query.all()
    assert len(units) == 3
    for unit in units:
        assert unit.upc == "000999000111"
        assert unit.sku == "SKU-KEEP"
        assert unit.description == "Full Attribute Coat"
        assert unit.style == "STYLE-Z"
        assert unit.color == "BURGUNDY"
        assert unit.size == "XXL"
        assert unit.location == "0080097422/0001"


def test_preview_is_compact(app, db, admin_user):
    celine, _, cel_ny, _, _ = _masters(db)
    rows = [_row(upc=f"{i:012d}", qty=1) for i in range(120)]
    preview = analyze(_xlsx(rows), "compact.xlsx", celine, cel_ny)
    assert preview["total_rows"] == 120
    assert preview["valid_rows"] == 120
    assert len(preview["rows"]) == 100
    rebuilt = preview_from_batch(ImportBatch.query.get(preview["batch_id"]))
    assert len(rebuilt["rows"]) == 100


def test_actual_celine_scale_import(app, db, admin_user):
    celine, _, cel_ny, _, _ = _masters(db)
    rows = _celine_rows()
    assert len(rows) == CELINE_ROWS
    assert sum(r[6] for r in rows) == CELINE_UNITS
    assert len({r[0] for r in rows}) == CELINE_UPCS
    assert len({r[3] for r in rows}) == CELINE_STYLES
    assert len({r[7] for r in rows}) == CELINE_LOCS
    source = _write_xlsx(rows)
    started = time.perf_counter()
    preview = analyze(source, "celine-scale.xlsx", celine, cel_ny)
    analyze_s = time.perf_counter() - started
    assert preview["has_blocking"] is False, preview["blocking"][:3]
    assert preview["total_rows"] == CELINE_ROWS
    assert preview["units"] == CELINE_UNITS
    assert preview["unique_upcs"] == CELINE_UPCS
    assert preview["unique_styles"] == CELINE_STYLES
    assert preview["unique_locations"] == CELINE_LOCS
    assert len(preview["rows"]) == 100
    started = time.perf_counter()
    batch = commit_import(preview, user=admin_user)
    process_s = time.perf_counter() - started
    assert batch.status == ImportBatchStatus.COMPLETED
    assert batch.units_created == CELINE_UNITS
    assert batch.transactions_created == CELINE_UNITS
    assert InventoryUnit.query.filter_by(import_batch_id=batch.id).count() == CELINE_UNITS
    assert (
        InventoryTransaction.query.filter_by(
            import_batch_id=batch.id, transaction_type=LedgerType.IMPORT
        ).count()
        == CELINE_UNITS
    )
    assert status_counts(client_id=celine.id, warehouse_id=cel_ny.id)["available"] == CELINE_UNITS
    assert_invariants(client_id=celine.id)
    _write_benchmark(
        {
            "actual_scale": {
                "rows": CELINE_ROWS,
                "units": CELINE_UNITS,
                "analyze_seconds": round(analyze_s, 3),
                "process_seconds": round(process_s, 3),
                "total_seconds": round(analyze_s + process_s, 3),
                "timings": preview.get("timings"),
                "message": batch.message,
            }
        }
    )


def test_capacity_25000_by_150000(app, db, admin_user):
    celine, _, cel_ny, _, _ = _masters(db)
    source = _write_xlsx(_capacity_rows())
    started = time.perf_counter()
    preview = analyze(source, "capacity.xlsx", celine, cel_ny)
    analyze_s = time.perf_counter() - started
    assert preview["has_blocking"] is False, preview["blocking"][:3]
    assert preview["total_rows"] == CAPACITY_ROWS
    assert preview["units"] == CAPACITY_UNITS
    started = time.perf_counter()
    batch = commit_import(preview, user=admin_user)
    process_s = time.perf_counter() - started
    assert batch.status == ImportBatchStatus.COMPLETED
    assert batch.units_created >= CAPACITY_UNITS
    assert batch.transactions_created == batch.units_created
    assert InventoryUnit.query.filter_by(import_batch_id=batch.id).count() == CAPACITY_UNITS
    assert status_counts(client_id=celine.id, warehouse_id=cel_ny.id)["available"] == CAPACITY_UNITS
    assert_invariants(client_id=celine.id)
    _write_benchmark(
        {
            "capacity": {
                "rows": CAPACITY_ROWS,
                "units": CAPACITY_UNITS,
                "analyze_seconds": round(analyze_s, 3),
                "process_seconds": round(process_s, 3),
                "total_seconds": round(analyze_s + process_s, 3),
                "timings": preview.get("timings"),
                "message": batch.message,
            }
        }
    )
