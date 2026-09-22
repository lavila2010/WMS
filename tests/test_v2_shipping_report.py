"""Shipping Report: single-sheet multi-order Excel export and tenant isolation."""

from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.constants import CartonStatus, OrderStatus, ShippingStatus
from app.extensions import db
from app.models import AuditEvent, Carton, CartonContent, InventoryUnit, Order, PickTicket
from app.services.allocation import allocate_order, approve_partial_allocation
from app.services.inventory_ledger import create_available_unit
from app.services.masters import create_division, map_division_warehouse
from app.services.pick_tickets import create_pick_ticket
from app.services.processing import (
    acquire_lock,
    close_order,
    close_order_short,
    ensure_open_carton,
    request_close,
    scan_upc,
    set_dimensions,
    set_weight,
)
from app.services.shipping import save_carton_tracking
from app.services.shipping_report import (
    ORDER_INFO_HEADERS,
    PRODUCT_HEADERS,
    WORKSHEET_NAME,
    ShippingReportAccessError,
    export_shipping_report,
    load_shipping_report_blocks,
    order_label,
)
from tests.conftest import create_user, form_data, login
from tests.sql_probe import count_sql
from tests.test_v2_phase03_orders import _line
from tests.test_v2_phase04_allocation import _import_order, _world


def _stock(client, warehouse, upc, qty, location="A-01", **attrs):
    for _ in range(qty):
        create_available_unit(
            client_id=client.id,
            warehouse_id=warehouse.id,
            upc=upc,
            location=location,
            sku=attrs.get("sku", "SKU-A"),
            description=attrs.get("description"),
            style=attrs.get("style"),
            color=attrs.get("color"),
            size=attrs.get("size"),
        )
    db.session.commit()


def _pack_scans(admin_user, order, scans, *, weight=6.8, dims=(18, 12, 8)):
    carton = ensure_open_carton(order, admin_user)
    set_dimensions(carton, *dims)
    db.session.commit()
    for upc in scans:
        scan_upc(order, carton, upc, admin_user)
    request_close(carton, admin_user)
    set_weight(carton, weight, user=admin_user)
    return Carton.query.get(carton.id)


def _ship_order(admin_user, order, carton_scans, trackings=None):
    acquire_lock(order, admin_user)
    cartons = []
    for item in carton_scans:
        cartons.append(_pack_scans(admin_user, order, item["scans"], weight=item.get("weight", 6.8), dims=item.get("dims", (18, 12, 8))))
    close_order(Order.query.get(order.id), admin_user)
    order = Order.query.get(order.id)
    cartons = Carton.query.filter_by(order_id=order.id).order_by(Carton.id).all()
    for carton, tracking in zip(cartons, trackings or []):
        if tracking:
            save_carton_tracking(
                order,
                carton,
                tracking_number=tracking.get("number", ""),
                carrier=tracking.get("carrier", ""),
                user=admin_user,
            )
    db.session.commit()
    return Order.query.get(order.id)


def _soho_order(db, admin_user, world=None):
    w = world or _world(db)
    attrs_a = dict(sku="BK001", description="Handbag Black", style="STYLE01", color="Black", size="OS")
    attrs_b = dict(sku="SL002", description="Leather Belt", style="STYLE02", color="Black", size="85")
    attrs_c = dict(sku="SH003", description="Sneaker White", style="STYLE03", color="White", size="39")
    _stock(w["celine"], w["cel_ny"], "3614270011111", 2, **attrs_a)
    _stock(w["celine"], w["cel_ny"], "3614270022222", 2, **attrs_b)
    _stock(w["celine"], w["cel_ny"], "3614270033333", 3, **attrs_c)
    order = _import_order(
        admin_user,
        w["celine"],
        w["cel_ecom"],
        [
            _line(order="100", upc="3614270011111", qty=2, customer="Soho NY"),
            _line(order="100", upc="3614270022222", qty=2, customer="Soho NY"),
            _line(order="100", upc="3614270033333", qty=3, customer="Soho NY"),
        ],
    )
    allocate_order(order, user=admin_user)
    create_pick_ticket(Order.query.get(order.id))
    closed = _ship_order(
        admin_user,
        Order.query.get(order.id),
        [
            {"scans": ["3614270011111", "3614270011111", "3614270022222", "3614270022222"], "weight": 6.8, "dims": (18, 12, 8)},
            {"scans": ["3614270033333", "3614270033333", "3614270033333"], "weight": 4.2, "dims": (14, 10, 6)},
        ],
        trackings=[
            {"number": "1Z999AAA001", "carrier": "UPS"},
            {"number": "1Z999AAA002", "carrier": "UPS"},
        ],
    )
    return w, closed


def _simple_closed(db, admin_user, world, number, customer, upc="UPC-A", qty=1, warehouse=None, division=None):
    warehouse = warehouse or world["cel_ny"]
    division = division or world["cel_ecom"]
    symbol = warehouse.warehouse_symbol
    _stock(world["celine"] if division.client_id == world["celine"].id else world["dior"], warehouse, upc, qty)
    client = world["celine"] if division.client_id == world["celine"].id else world["dior"]
    order = _import_order(
        admin_user,
        client,
        division,
        [_line(order=str(number), warehouse=symbol, upc=upc, qty=qty, customer=customer)],
    )
    allocate_order(order, user=admin_user)
    create_pick_ticket(Order.query.get(order.id))
    return _ship_order(
        admin_user,
        Order.query.get(order.id),
        [{"scans": [upc] * qty, "weight": 2.2, "dims": (10, 8, 6)}],
        trackings=[{"number": f"1Z{number}TRACK", "carrier": "UPS"}],
    )


def _workbook(data):
    return load_workbook(BytesIO(data))


def _sheet_values(ws):
    return [[cell.value for cell in row] for row in ws.iter_rows()]


def _find_row(ws, text):
    for row in ws.iter_rows(min_col=1, max_col=1):
        if row[0].value == text:
            return row[0].row
    raise AssertionError(f"Row not found: {text}")


def _export(client, order_ids, path="/"):
    payload = form_data(client, {}, path=path)
    payload["order_ids"] = [str(oid) for oid in order_ids]
    return client.post("/orders/shipping-report/export", data=payload)


def test_01_route_loads(app, db, admin_user, admin_client):
    resp = admin_client.get("/orders/shipping-report")
    html = resp.get_data(as_text=True)
    assert resp.status_code == 200
    assert "Shipping Report" in html
    assert "Select All Visible" in html
    assert "Export Selected Orders" in html
    assert "Clear Selection" in html
    assert 'name="order_status"' in html
    assert 'name="shipping_status"' in html
    assert 'name="carrier"' in html
    assert 'name="created_from"' in html
    assert 'name="created_to"' in html
    assert 'name="closed_from"' in html
    assert 'name="closed_to"' in html
    assert 'name="q"' in html
    assert "Ordered Units" in html
    assert "Tracking Summary" in html
    assert 'option value="50" selected' in html
    assert "25" in html and "100" in html


def test_02_03_permissions_enforced(app, db, admin_user):
    w = _world(db)
    create_user("shipview", perms=["SHIPPING_VIEW", "ORDERS_VIEW", "DASHBOARD_VIEW"], clients=[w["celine"].id])
    create_user(
        "reportview",
        perms=["SHIPPING_REPORT_VIEW", "ORDERS_VIEW", "DASHBOARD_VIEW"],
        clients=[w["celine"].id],
    )
    create_user(
        "reportexport",
        perms=["SHIPPING_REPORT_VIEW", "SHIPPING_REPORT_EXPORT", "DASHBOARD_VIEW"],
        clients=[w["celine"].id],
    )
    view_only = app.test_client()
    login(view_only, "shipview")
    assert view_only.get("/orders/shipping-report").status_code == 403
    assert _export(view_only, [1]).status_code == 403

    report_view = app.test_client()
    login(report_view, "reportview")
    assert report_view.get("/orders/shipping-report").status_code == 200
    html = report_view.get("/orders/shipping-report").get_data(as_text=True)
    assert "Export Selected Orders" not in html
    assert _export(report_view, [1]).status_code == 403

    exporter = app.test_client()
    login(exporter, "reportexport")
    assert exporter.get("/orders/shipping-report").status_code == 200
    assert _export(exporter, []).status_code in {302, 200}


def test_04_05_authorized_orders_and_hidden_clients(app, db, admin_user, admin_client):
    w = _world(db)
    cel = _simple_closed(db, admin_user, w, "100", "Soho NY")
    dio = _simple_closed(db, admin_user, w, "DIO-1", "Paris", upc="UPC-D", division=w["dio_ecom"], warehouse=w["dio_ny"])
    html = admin_client.get("/orders/shipping-report").get_data(as_text=True)
    assert "CEL - 100 - Soho NY" in html
    assert "DIO - DIO-1 - Paris" in html
    create_user(
        "cel-user",
        perms=["SHIPPING_REPORT_VIEW", "SHIPPING_REPORT_EXPORT"],
        clients=[w["celine"].id],
    )
    limited = app.test_client()
    login(limited, "cel-user")
    html = limited.get("/orders/shipping-report").get_data(as_text=True)
    assert "CEL - 100 - Soho NY" in html
    assert "DIO - DIO-1 - Paris" not in html
    assert "Paris" not in html


def test_06_07_08_filters(app, db, admin_user, admin_client):
    w = _world(db)
    retail = create_division(w["celine"], "Retail", "RTL")
    map_division_warehouse(retail, w["cel_ny"])
    db.session.commit()
    ecom = _simple_closed(db, admin_user, w, "100", "Soho NY")
    rtl = _simple_closed(db, admin_user, w, "200", "Retail Store", upc="UPC-R", division=retail)
    nj = _simple_closed(db, admin_user, w, "300", "New Jersey", upc="UPC-N", warehouse=w["cel_nj"])
    dio = _simple_closed(db, admin_user, w, "DIO-1", "Paris", upc="UPC-D", division=w["dio_ecom"], warehouse=w["dio_ny"])

    client_html = admin_client.get(f"/orders/shipping-report?client_id={w['celine'].id}").get_data(as_text=True)
    assert "CEL - 100 - Soho NY" in client_html
    assert "Paris" not in client_html

    div_html = admin_client.get(f"/orders/shipping-report?division_id={retail.id}").get_data(as_text=True)
    assert "CEL - 200 - Retail Store" in div_html
    assert "Soho NY" not in div_html

    wh_html = admin_client.get(f"/orders/shipping-report?warehouse_id={w['cel_nj'].id}").get_data(as_text=True)
    assert "CEL - 300 - New Jersey" in wh_html
    assert "Soho NY" not in wh_html


def test_09_10_11_12_search(app, db, admin_user, admin_client):
    w, order = _soho_order(db, admin_user)
    ticket = PickTicket.query.filter_by(order_id=order.id).one()
    label = "CEL - 100 - Soho NY"
    assert label in admin_client.get(f"/orders/shipping-report?q={order.wms_order_id}").get_data(as_text=True)
    assert label in admin_client.get("/orders/shipping-report?q=100").get_data(as_text=True)
    assert label in admin_client.get("/orders/shipping-report?q=Soho").get_data(as_text=True)
    assert label in admin_client.get("/orders/shipping-report?q=1Z999AAA001").get_data(as_text=True)
    assert ticket.pick_ticket_number
    assert label in admin_client.get(f"/orders/shipping-report?q={ticket.pick_ticket_number}").get_data(as_text=True)
    miss = admin_client.get("/orders/shipping-report?q=NO-SUCH-TRACK").get_data(as_text=True)
    assert label not in miss


def test_13_14_multiselect_and_select_all_visible(app, db, admin_user, admin_client):
    w = _world(db)
    first = _simple_closed(db, admin_user, w, "100", "Soho NY")
    second = _simple_closed(db, admin_user, w, "101", "Madison Ave", upc="UPC-B")
    html = admin_client.get("/orders/shipping-report").get_data(as_text=True)
    assert 'id="select-all-visible"' in html
    assert html.count('class="order-check"') >= 2
    assert "order-check" in html
    resp = _export(admin_client, [first.id, second.id])
    assert resp.status_code == 200
    wb = _workbook(resp.data)
    ws = wb.active
    labels = [row[0].value for row in ws.iter_rows(min_col=1, max_col=1) if row[0].value]
    assert "CEL - 100 - Soho NY" in labels
    assert "CEL - 101 - Madison Ave" in labels


def test_15_16_17_18_19_20_21_22_export_layout(app, db, admin_user, admin_client):
    w, soho = _soho_order(db, admin_user)
    madison = _simple_closed(db, admin_user, w, "101", "Madison Ave", upc="UPC-B")
    one = _export(admin_client, [soho.id])
    assert one.status_code == 200
    assert "Shipping_Report_" in (one.headers.get("Content-Disposition") or "")
    assert one.data[:2] == b"PK"
    wb = _workbook(one.data)
    assert wb.sheetnames == [WORKSHEET_NAME]
    assert len(wb.worksheets) == 1
    ws = wb[WORKSHEET_NAME]
    assert ws.title == "Shipping Report"

    many = _export(admin_client, [soho.id, madison.id])
    wb = _workbook(many.data)
    ws = wb.active
    assert wb.sheetnames == ["Shipping Report"]
    soho_row = _find_row(ws, "CEL - 100 - Soho NY")
    madison_row = _find_row(ws, "CEL - 101 - Madison Ave")
    assert soho_row < madison_row
    headers = [ws.cell(soho_row + 1, col).value for col in range(1, 17)]
    values = [ws.cell(soho_row + 2, col).value for col in range(1, 17)]
    assert headers == ORDER_INFO_HEADERS
    assert values[0] == "CEL"
    assert values[1] == "100"
    assert values[2] == "Soho NY"
    assert values[3] == soho.wms_order_id
    assert "Client:" not in "".join(str(v) for v in values)
    assert ws.cell(soho_row, 1).value == "CEL - 100 - Soho NY"
    assert ws.merge_cells  # merged label exists
    assert order_label("CEL", "100", "Soho NY") == "CEL - 100 - Soho NY"
    assert order_label("CEL", "100", "") == "CEL - 100"
    assert order_label("CEL", "100", None) == "CEL - 100"


def test_23_24_25_26_carton_tracking_and_tickets(app, db, admin_user, admin_client):
    w, order = _soho_order(db, admin_user)
    ticket = PickTicket.query.filter_by(order_id=order.id).one()
    cartons = Carton.query.filter_by(order_id=order.id).order_by(Carton.id).all()
    data, _, _ = export_shipping_report(admin_user, [order.id])
    ws = _workbook(data).active
    text = "\n".join(str(row[0].value or "") for row in ws.iter_rows(min_col=1, max_col=1))
    assert cartons[0].carton_number in text
    assert cartons[1].carton_number in text
    assert "1Z999AAA001" in text
    assert "1Z999AAA002" in text
    first = text.index("1Z999AAA001")
    second = text.index("1Z999AAA002")
    assert first < second
    assert ticket.pick_ticket_number in text
    assert f"Carton {cartons[0].carton_number}" in text
    assert f"Pick Ticket: {ticket.pick_ticket_number}" in text


def test_27_multiple_pick_ticket_waves(app, db, admin_user):
    w = _world(db)
    _stock(w["celine"], w["cel_ny"], "UPC-A", 1)
    order = _import_order(admin_user, w["celine"], w["cel_ecom"], [_line(order="8801", upc="UPC-A", qty=2, customer="Wave Co")])
    allocate_order(order, user=admin_user)
    approve_partial_allocation(Order.query.get(order.id), user=admin_user)
    ticket1 = create_pick_ticket(Order.query.get(order.id))
    acquire_lock(Order.query.get(order.id), admin_user)
    _pack_scans(admin_user, Order.query.get(order.id), ["UPC-A"], weight=1.5, dims=(8, 6, 4))
    close_order(Order.query.get(order.id), admin_user)
    _stock(w["celine"], w["cel_ny"], "UPC-A", 1, location="B-02")
    allocate_order(Order.query.get(order.id), user=admin_user)
    ticket2 = create_pick_ticket(Order.query.get(order.id))
    assert ticket2.id != ticket1.id
    acquire_lock(Order.query.get(order.id), admin_user)
    _pack_scans(admin_user, Order.query.get(order.id), ["UPC-A"], weight=1.6, dims=(8, 6, 4))
    close_order(Order.query.get(order.id), admin_user)
    order = Order.query.get(order.id)
    cartons = Carton.query.filter_by(order_id=order.id).order_by(Carton.id).all()
    assert cartons[0].pick_ticket_id == ticket1.id
    assert cartons[1].pick_ticket_id == ticket2.id
    data, _, _ = export_shipping_report(admin_user, [order.id])
    text = "\n".join(str(row[0].value or "") for row in _workbook(data).active.iter_rows(min_col=1, max_col=1))
    assert ticket1.pick_ticket_number in text
    assert ticket2.pick_ticket_number in text
    assert text.index(ticket1.pick_ticket_number) < text.index(ticket2.pick_ticket_number)


def test_28_35_products_from_carton_contents(app, db, admin_user):
    _, order = _soho_order(db, admin_user)
    data, _, stats = export_shipping_report(admin_user, [order.id])
    ws = _workbook(data).active
    rows = _sheet_values(ws)
    product_headers = next(row for row in rows if row[:7] == PRODUCT_HEADERS)
    assert product_headers[:7] == PRODUCT_HEADERS
    products = [row[:7] for row in rows if row[0] in {"3614270011111", "3614270022222", "3614270033333"}]
    assert ["3614270011111", "BK001", "Handbag Black", "STYLE01", "Black", "OS", 2] in products
    assert ["3614270022222", "SL002", "Leather Belt", "STYLE02", "Black", "85", 2] in products
    assert ["3614270033333", "SH003", "Sneaker White", "STYLE03", "White", "39", 3] in products
    assert stats["product_rows"] == 3
    first_upc_rows = [index for index, row in enumerate(rows) if row[0] == "3614270011111"]
    assert len(first_upc_rows) == 1
    sneaker_row = next(index for index, row in enumerate(rows) if row[0] == "3614270033333")
    assert sneaker_row > first_upc_rows[0]


def test_29_30_qty_aggregation_and_separation(app, db, admin_user):
    w = _world(db)
    _stock(w["celine"], w["cel_ny"], "UPC111", 3, sku="SKU111", description="Same UPC")
    order = _import_order(admin_user, w["celine"], w["cel_ecom"], [_line(order="777", upc="UPC111", qty=3, customer="Agg")])
    allocate_order(order, user=admin_user)
    create_pick_ticket(Order.query.get(order.id))
    closed = _ship_order(
        admin_user,
        Order.query.get(order.id),
        [
            {"scans": ["UPC111", "UPC111"], "weight": 3.0, "dims": (10, 8, 6)},
            {"scans": ["UPC111"], "weight": 1.5, "dims": (8, 6, 4)},
        ],
        trackings=[{"number": "TRK-A", "carrier": "UPS"}, {"number": "TRK-B", "carrier": "FEDEX"}],
    )
    data, _, _ = export_shipping_report(admin_user, [closed.id])
    rows = _sheet_values(_workbook(data).active)
    upc_rows = [row for row in rows if row[0] == "UPC111"]
    assert [row[6] for row in upc_rows] == [2, 1]


def test_36_37_closed_short_excludes_missing(app, db, admin_user):
    w = _world(db)
    _stock(w["celine"], w["cel_ny"], "UPC-A", 2)
    order = _import_order(admin_user, w["celine"], w["cel_ecom"], [_line(order="15", upc="UPC-A", qty=2, customer="Short Store")])
    allocate_order(order, user=admin_user)
    create_pick_ticket(Order.query.get(order.id))
    acquire_lock(Order.query.get(order.id), admin_user)
    _pack_scans(admin_user, Order.query.get(order.id), ["UPC-A"], weight=2.0, dims=(10, 8, 6))
    close_order_short(Order.query.get(order.id), admin_user, confirmed=True)
    order = Order.query.get(order.id)
    assert order.short_closed is True
    assert sum(line.qty_ordered for line in order.lines) == 2
    assert sum(line.qty_shipped for line in order.lines) == 1
    assert sum(line.qty_short for line in order.lines) == 1
    assert InventoryUnit.query.filter_by(status="MISSING").count() == 1
    data, _, stats = export_shipping_report(admin_user, [order.id])
    ws = _workbook(data).active
    label_row = _find_row(ws, "CEL - 15 - Short Store")
    values = [ws.cell(label_row + 2, col).value for col in range(1, 17)]
    assert values[10] == 2
    assert values[11] == 1
    assert values[12] == 1
    products = [row for row in _sheet_values(ws) if row[0] == "UPC-A"]
    assert len(products) == 1
    assert products[0][6] == 1
    assert stats["product_rows"] == 1
    assert CartonContent.query.filter_by(carton_id=Carton.query.filter_by(order_id=order.id).one().id).count() == 1


def test_38_order_without_cartons(app, db, admin_user):
    w = _world(db)
    order = _import_order(admin_user, w["celine"], w["cel_ecom"], [_line(order="900", customer="No Boxes")])
    data, _, stats = export_shipping_report(admin_user, [order.id])
    text = "\n".join(str(row[0].value or "") for row in _workbook(data).active.iter_rows(min_col=1, max_col=1))
    assert "CEL - 900 - No Boxes" in text
    assert "No cartons available." in text
    assert stats["cartons"] == 0


def test_39_carton_without_tracking(app, db, admin_user):
    w = _world(db)
    _stock(w["celine"], w["cel_ny"], "UPC-A", 1)
    order = _import_order(admin_user, w["celine"], w["cel_ecom"], [_line(order="901", upc="UPC-A", qty=1, customer="No Track")])
    allocate_order(order, user=admin_user)
    create_pick_ticket(Order.query.get(order.id))
    closed = _ship_order(admin_user, Order.query.get(order.id), [{"scans": ["UPC-A"]}], trackings=[])
    data, _, _ = export_shipping_report(admin_user, [closed.id])
    text = "\n".join(str(row[0].value or "") for row in _workbook(data).active.iter_rows(min_col=1, max_col=1))
    assert "Tracking: " in text
    assert "1Z" not in text
    assert "N/A" not in text
    assert "TRACK-PENDING" not in text


def test_40_carton_without_products(app, db, admin_user):
    w = _world(db)
    order = _import_order(admin_user, w["celine"], w["cel_ecom"], [_line(order="902", customer="Empty Carton")])
    carton = Carton(
        client_id=order.client_id,
        order_id=order.id,
        carton_number="902-EMPTY",
        status=CartonStatus.CLOSED,
        weight=1.0,
        weight_unit="lb",
        length=6,
        width=6,
        height=6,
        dimension_unit="in",
    )
    db.session.add(carton)
    db.session.commit()
    data, _, _ = export_shipping_report(admin_user, [order.id])
    text = "\n".join(str(row[0].value or "") for row in _workbook(data).active.iter_rows(min_col=1, max_col=1))
    assert "Carton 902-EMPTY" in text
    assert "No products" in text


def test_41_cross_client_order_ids_rejected(app, db, admin_user):
    w = _world(db)
    cel = _simple_closed(db, admin_user, w, "100", "Soho NY")
    dio = _simple_closed(db, admin_user, w, "DIO-1", "Paris", upc="UPC-D", division=w["dio_ecom"], warehouse=w["dio_ny"])
    create_user(
        "cel-export",
        perms=["SHIPPING_REPORT_VIEW", "SHIPPING_REPORT_EXPORT"],
        clients=[w["celine"].id],
    )
    limited = app.test_client()
    login(limited, "cel-export")
    assert _export(limited, [cel.id]).status_code == 200
    assert _export(limited, [dio.id]).status_code == 404
    assert _export(limited, [cel.id, dio.id]).status_code == 404


def test_41b_cross_client_service_user(app, db, admin_user):
    w = _world(db)
    dio = _simple_closed(db, admin_user, w, "DIO-1", "Paris", upc="UPC-D", division=w["dio_ecom"], warehouse=w["dio_ny"])
    user = create_user(
        "cel-svc",
        perms=["SHIPPING_REPORT_VIEW", "SHIPPING_REPORT_EXPORT"],
        clients=[w["celine"].id],
    )
    db.session.refresh(user)
    with pytest.raises(ShippingReportAccessError):
        export_shipping_report(user, [dio.id])
    with pytest.raises(ShippingReportAccessError):
        export_shipping_report(user, [999999])


def test_42_export_audit(app, db, admin_user):
    _, order = _soho_order(db, admin_user)
    _, _, stats = export_shipping_report(admin_user, [order.id])
    db.session.commit()
    event = AuditEvent.query.filter_by(event_type="SHIPPING_REPORT_EXPORTED").order_by(AuditEvent.id.desc()).first()
    assert event is not None
    assert f"selected_orders={stats['selected_orders']}" in event.detail
    assert f"cartons={stats['cartons']}" in event.detail
    assert f"product_rows={stats['product_rows']}" in event.detail
    assert "Soho" not in (event.detail or "")
    assert "address" not in (event.detail or "").lower()


def test_43_no_n_plus_one(app, db, admin_user):
    w = _world(db)
    first = _simple_closed(db, admin_user, w, "100", "One")
    extras = [_simple_closed(db, admin_user, w, str(200 + i), f"Cust {i}", upc=f"UPC-{i}") for i in range(4)]
    with count_sql() as one:
        load_shipping_report_blocks(admin_user, [first.id])
    with count_sql() as many:
        load_shipping_report_blocks(admin_user, [first.id, *[row.id for row in extras]])
    assert one["count"] <= 20
    assert many["count"] <= one["count"] + 3
    assert many["count"] < one["count"] * 2


def test_44_openpyxl_reopen(app, db, admin_user, admin_client):
    _, order = _soho_order(db, admin_user)
    resp = _export(admin_client, [order.id])
    wb = load_workbook(BytesIO(resp.data))
    assert wb[WORKSHEET_NAME]["A1"].value == "CEL - 100 - Soho NY"
    again = load_workbook(BytesIO(resp.data))
    assert again.sheetnames == ["Shipping Report"]


def test_newest_first_and_blank_customer_label(app, db, admin_user, admin_client):
    w = _world(db)
    older = _simple_closed(db, admin_user, w, "100", "Older")
    newer = _import_order(admin_user, w["celine"], w["cel_ecom"], [_line(order="101", customer="Placeholder")])
    newer.customer = ""
    db.session.commit()
    html = admin_client.get("/orders/shipping-report").get_data(as_text=True)
    assert html.index("CEL - 101") < html.index("CEL - 100 - Older")
    assert "CEL - 101" in html
    assert "CEL - 101 -" not in html
    data, _, _ = export_shipping_report(admin_user, [newer.id])
    assert _workbook(data).active["A1"].value == "CEL - 101"
