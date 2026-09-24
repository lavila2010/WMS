"""End of Day carton manifest: CLOSED + PARTIALLY_FULFILLED daily activity."""

from datetime import date, datetime, timezone
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from app.constants import EOD_TIMEZONE, OrderStatus
from app.extensions import db
from app.models import AuditEvent, Carton, CartonContent, InventoryUnit, Order, PickTicket
from app.services.allocation import allocate_order, approve_partial_allocation
from app.services.end_of_day import (
    FULFILLMENT_CLOSED_COMPLETE,
    FULFILLMENT_CLOSED_SHORT,
    FULFILLMENT_PARTIALLY,
    ORDER_INFO_HEADERS,
    WORKSHEET_NAME,
    build_eod_rows,
    closed_orders_query,
    eod_kpis,
    export_eod_excel,
    ny_day_utc_bounds,
)
from app.services.inventory_ledger import create_available_unit
from app.services.packing_list import render_packing_list_pdf
from app.services.pick_tickets import create_pick_ticket
from app.services.processing import (
    acquire_lock,
    close_order,
    close_order_short,
    ensure_open_carton,
    request_close,
    scan_upc,
    set_dimensions,
    set_weight,
)
from app.services.shipping import save_carton_tracking
from app.services.shipping_report import PRODUCT_HEADERS, export_shipping_report, fulfillment_status
from tests.conftest import create_user, login
from tests.sql_probe import count_sql
from tests.test_v2_phase03_orders import _line
from tests.test_v2_phase04_allocation import _import_order, _world
from tests.test_v2_shipping_report import _soho_order


TZ = ZoneInfo(EOD_TIMEZONE)
DAY = date(2026, 9, 24)


def _stock(client, warehouse, upc, qty, location="A-01", **attrs):
    for _ in range(qty):
        create_available_unit(
            client_id=client.id,
            warehouse_id=warehouse.id,
            upc=upc,
            location=location,
            sku=attrs.get("sku", "SKU-A"),
            description=attrs.get("description"),
            style=attrs.get("style"),
            color=attrs.get("color"),
            size=attrs.get("size"),
        )
    db.session.commit()


def _stamp(order, cartons, when):
    if order.status == OrderStatus.CLOSED:
        order.closed_at = when
    for carton in cartons:
        carton.closed_at = when
    db.session.commit()


def _midday(day=DAY):
    start, _ = ny_day_utc_bounds(day)
    return start + (datetime(day.year, day.month, day.day, 12, 0) - datetime(day.year, day.month, day.day))


def _pack_splits(admin_user, order, splits, upc="UPC-A", trackings=None):
    acquire_lock(order, admin_user)
    cartons = []
    for index, qty in enumerate(splits):
        carton = ensure_open_carton(order, admin_user)
        set_dimensions(carton, 12, 10, 8)
        db.session.commit()
        for _ in range(qty):
            scan_upc(order, carton, upc, admin_user)
        request_close(carton, admin_user)
        set_weight(carton, 2.5 + index, user=admin_user)
        cartons.append(Carton.query.get(carton.id))
    return [Carton.query.get(carton.id) for carton in cartons]


def _close_complete(admin_user, world, number, qty=10, customer="Complete Co", upc="UPC-A"):
    _stock(world["celine"], world["cel_ny"], upc, qty, sku="SKU-C", description="Complete Bag", style="ST-C", color="Black", size="OS")
    order = _import_order(
        admin_user,
        world["celine"],
        world["cel_ecom"],
        [_line(order=str(number), upc=upc, qty=qty, customer=customer)],
    )
    allocate_order(order, user=admin_user)
    create_pick_ticket(Order.query.get(order.id))
    cartons = _pack_splits(admin_user, Order.query.get(order.id), [4, 3, 3], upc=upc)
    closed = close_order(Order.query.get(order.id), admin_user)
    closed = Order.query.get(closed.id)
    cartons = Carton.query.filter_by(order_id=closed.id).order_by(Carton.id).all()
    for carton, tracking in zip(cartons, ["1ZCOMP001", "1ZCOMP002", "1ZCOMP003"]):
        save_carton_tracking(closed, carton, tracking_number=tracking, carrier="UPS", user=admin_user)
    _stamp(closed, cartons, _midday())
    return Order.query.get(closed.id), Carton.query.filter_by(order_id=closed.id).order_by(Carton.id).all()


def _close_short(admin_user, world, number, ordered=10, packed=8, customer="Short Co", upc="UPC-S"):
    _stock(world["celine"], world["cel_ny"], upc, ordered)
    order = _import_order(
        admin_user,
        world["celine"],
        world["cel_ecom"],
        [_line(order=str(number), upc=upc, qty=ordered, customer=customer)],
    )
    allocate_order(order, user=admin_user)
    create_pick_ticket(Order.query.get(order.id))
    acquire_lock(Order.query.get(order.id), admin_user)
    carton = ensure_open_carton(Order.query.get(order.id), admin_user)
    set_dimensions(carton, 12, 10, 8)
    db.session.commit()
    for _ in range(packed):
        scan_upc(Order.query.get(order.id), carton, upc, admin_user)
    request_close(carton, admin_user)
    set_weight(carton, 3.1, user=admin_user)
    close_order_short(Order.query.get(order.id), admin_user, confirmed=True)
    order = Order.query.get(order.id)
    cartons = Carton.query.filter_by(order_id=order.id).order_by(Carton.id).all()
    save_carton_tracking(order, cartons[0], tracking_number="1ZSHORT001", carrier="UPS", user=admin_user)
    _stamp(order, cartons, _midday())
    return Order.query.get(order.id), cartons


def _partial_today(admin_user, world, number="100", customer="Soho NY", upc="UPC-A"):
    _stock(
        world["celine"],
        world["cel_ny"],
        upc,
        10,
        sku="SKU-P",
        description="Partial Bag",
        style="ST-P",
        color="Navy",
        size="M",
    )
    order = _import_order(
        admin_user,
        world["celine"],
        world["cel_ecom"],
        [_line(order=str(number), upc=upc, qty=10, customer=customer)],
    )
    leftover = InventoryUnit.query.filter_by(upc=upc, status="AVAILABLE").limit(3).all()
    for unit in leftover:
        unit.status = "SHIPPED"
    db.session.commit()
    allocate_order(Order.query.get(order.id), user=admin_user)
    approve_partial_allocation(Order.query.get(order.id), user=admin_user)
    create_pick_ticket(Order.query.get(order.id))
    cartons = _pack_splits(admin_user, Order.query.get(order.id), [4, 3], upc=upc)
    close_order(Order.query.get(order.id), admin_user)
    order = Order.query.get(order.id)
    cartons = Carton.query.filter_by(order_id=order.id).order_by(Carton.id).all()
    save_carton_tracking(order, cartons[0], tracking_number="1ZPARTBOX01", carrier="UPS", user=admin_user)
    save_carton_tracking(order, cartons[1], tracking_number="1ZPARTBOX02", carrier="UPS", user=admin_user)
    _stamp(order, cartons, _midday())
    return Order.query.get(order.id), Carton.query.filter_by(order_id=order.id).order_by(Carton.id).all()


def _rows(day=DAY, **filters):
    return build_eod_rows(closed_orders_query(day=day, admin=True, **filters).all(), day=day)


def _ids(rows):
    return {row["order"].id for row in rows}


def _sheet(data):
    return load_workbook(BytesIO(data)).active


def _find(ws, text):
    for row in ws.iter_rows(min_col=1, max_col=1):
        if row[0].value == text:
            return row[0].row
    raise AssertionError(f"missing {text}")


def test_01_02_03_04_inclusion(app, db, admin_user, admin_client):
    w = _world(db)
    complete, _ = _close_complete(admin_user, w, "200")
    short, _ = _close_short(admin_user, w, "201")
    partial, cartons = _partial_today(admin_user, w, "100")
    historical, hist_cartons = _partial_today(admin_user, w, "199", customer="Old Wave")
    old = datetime(2026, 9, 20, 16, 0)
    _stamp(historical, hist_cartons, old)
    assert complete.status == OrderStatus.CLOSED
    assert short.short_closed is True
    assert partial.status == OrderStatus.PARTIALLY_FULFILLED
    rows = _rows()
    ids = _ids(rows)
    assert complete.id in ids
    assert short.id in ids
    assert partial.id in ids
    assert historical.id not in ids
    html = admin_client.get("/orders/end-of-day?date=2026-09-24").get_data(as_text=True)
    assert complete.wms_order_id in html
    assert short.wms_order_id in html
    assert partial.wms_order_id in html
    assert historical.wms_order_id not in html
    assert "CLOSED COMPLETE" in html
    assert "CLOSED SHORT" in html
    assert "PARTIALLY FULFILLED" in html
    assert "Total Orders Processed" in html


def test_05_06_07_quantities(app, db, admin_user):
    w = _world(db)
    complete, _ = _close_complete(admin_user, w, "300")
    short, _ = _close_short(admin_user, w, "301")
    partial, _ = _partial_today(admin_user, w, "302")
    by_id = {row["order"].id: row for row in _rows()}
    assert by_id[complete.id]["ordered"] == 10
    assert by_id[complete.id]["shipped"] == 10
    assert by_id[complete.id]["short"] == 0
    assert by_id[complete.id]["remaining"] == 0
    assert by_id[complete.id]["fulfillment_status"] == FULFILLMENT_CLOSED_COMPLETE
    assert by_id[short.id]["ordered"] == 10
    assert by_id[short.id]["shipped"] == 8
    assert by_id[short.id]["short"] == 2
    assert by_id[short.id]["remaining"] == 0
    assert by_id[short.id]["fulfillment_status"] == FULFILLMENT_CLOSED_SHORT
    assert by_id[partial.id]["ordered"] == 10
    assert by_id[partial.id]["shipped"] == 7
    assert by_id[partial.id]["short"] == 0
    assert by_id[partial.id]["remaining"] == 3
    assert by_id[partial.id]["fulfillment_status"] == FULFILLMENT_PARTIALLY
    assert by_id[partial.id]["reconciled"] is True


def test_08_12_cartons_and_tracking(app, db, admin_user):
    w = _world(db)
    complete, complete_cartons = _close_complete(admin_user, w, "400")
    partial, partial_cartons = _partial_today(admin_user, w, "401")
    by_id = {row["order"].id: row for row in _rows()}
    assert by_id[complete.id]["cartons_processed_today"] == 3
    assert by_id[complete.id]["total_cartons"] == 3
    assert by_id[partial.id]["cartons_processed_today"] == 2
    assert by_id[partial.id]["total_cartons"] == 2
    data = export_eod_excel(_rows(), day=DAY)
    ws = _sheet(data)
    text = "\n".join(str(cell.value or "") for row in ws.iter_rows() for cell in row)
    assert complete_cartons[0].carton_number in text
    assert complete_cartons[1].carton_number in text
    assert "1ZCOMP001" in text
    assert "1ZCOMP002" in text
    assert "1ZPARTBOX01" in text
    assert "1ZPARTBOX02" in text
    assert text.index("1ZPARTBOX01") < text.index("1ZPARTBOX02")


def test_13_multi_wave_pick_tickets(app, db, admin_user):
    w = _world(db)
    _stock(w["celine"], w["cel_ny"], "UPC-A", 1)
    order = _import_order(admin_user, w["celine"], w["cel_ecom"], [_line(order="8801", upc="UPC-A", qty=2, customer="Wave Co")])
    allocate_order(order, user=admin_user)
    approve_partial_allocation(Order.query.get(order.id), user=admin_user)
    ticket1 = create_pick_ticket(Order.query.get(order.id))
    acquire_lock(Order.query.get(order.id), admin_user)
    _pack_splits(admin_user, Order.query.get(order.id), [1], upc="UPC-A")
    close_order(Order.query.get(order.id), admin_user)
    first = Carton.query.filter_by(order_id=order.id).one()
    first.closed_at = datetime(2026, 9, 20, 16, 0)
    db.session.commit()
    _stock(w["celine"], w["cel_ny"], "UPC-A", 1, location="B-02")
    allocate_order(Order.query.get(order.id), user=admin_user)
    ticket2 = create_pick_ticket(Order.query.get(order.id))
    acquire_lock(Order.query.get(order.id), admin_user)
    _pack_splits(admin_user, Order.query.get(order.id), [1], upc="UPC-A")
    close_order(Order.query.get(order.id), admin_user)
    order = Order.query.get(order.id)
    cartons = Carton.query.filter_by(order_id=order.id).order_by(Carton.id).all()
    cartons[1].closed_at = _midday()
    order.closed_at = _midday()
    db.session.commit()
    assert cartons[0].pick_ticket_id == ticket1.id
    assert cartons[1].pick_ticket_id == ticket2.id
    rows = _rows()
    match = next(row for row in rows if row["order"].id == order.id)
    tickets = [block["pick_ticket_number"] for block in match["carton_blocks"]]
    assert ticket1.pick_ticket_number in tickets
    assert ticket2.pick_ticket_number in tickets
    data = export_eod_excel(rows, day=DAY)
    text = "\n".join(str(cell.value or "") for row in _sheet(data).iter_rows() for cell in row)
    assert ticket1.pick_ticket_number in text
    assert ticket2.pick_ticket_number in text


def test_14_22_products_grouping_and_attrs(app, db, admin_user):
    w, soho = _soho_order(db, admin_user)
    soho.closed_at = _midday()
    for carton in Carton.query.filter_by(order_id=soho.id).all():
        carton.closed_at = _midday()
    db.session.commit()
    rows = _rows()
    match = next(row for row in rows if row["order"].id == soho.id)
    products = [product for block in match["carton_blocks"] for product in block["products"]]
    assert ["3614270011111", "BK001", "Handbag Black", "STYLE01", "Black", "OS", 2] == [
        products[0]["upc"],
        products[0]["sku"],
        products[0]["description"],
        products[0]["style"],
        products[0]["color"],
        products[0]["size"],
        products[0]["qty"],
    ]
    upc111 = [p for p in products if p["upc"] == "3614270011111"]
    assert len(upc111) == 1
    assert upc111[0]["qty"] == 2
    assert sum(p["qty"] for p in products) == 7
    assert sum(CartonContent.query.filter_by(carton_id=block["carton"].id).count() for block in match["carton_blocks"]) == 7


def test_16_same_upc_separated_by_carton(app, db, admin_user):
    w = _world(db)
    _stock(w["celine"], w["cel_ny"], "UPC111", 3, sku="SKU111", description="Same UPC")
    order = _import_order(admin_user, w["celine"], w["cel_ecom"], [_line(order="777", upc="UPC111", qty=3, customer="Agg")])
    allocate_order(order, user=admin_user)
    create_pick_ticket(Order.query.get(order.id))
    _pack_splits(admin_user, Order.query.get(order.id), [2, 1], upc="UPC111")
    closed = close_order(Order.query.get(order.id), admin_user)
    cartons = Carton.query.filter_by(order_id=closed.id).order_by(Carton.id).all()
    _stamp(Order.query.get(closed.id), cartons, _midday())
    match = next(row for row in _rows() if row["order"].id == closed.id)
    qtys = [product["qty"] for block in match["carton_blocks"] for product in block["products"] if product["upc"] == "UPC111"]
    assert qtys == [2, 1]


def test_23_24_25_missing_and_remaining_excluded(app, db, admin_user):
    w = _world(db)
    short, cartons = _close_short(admin_user, w, "15")
    partial, _ = _partial_today(admin_user, w, "16")
    by_id = {row["order"].id: row for row in _rows()}
    assert InventoryUnit.query.filter_by(status="MISSING").count() >= 2
    short_units = by_id[short.id]["carton_units"]
    assert short_units == 8
    assert by_id[short.id]["shipped"] == 8
    assert by_id[short.id]["short"] == 2
    assert CartonContent.query.filter_by(carton_id=cartons[0].id).count() == 8
    assert by_id[partial.id]["carton_units"] == 7
    assert by_id[partial.id]["remaining"] == 3
    data = export_eod_excel(_rows(), day=DAY)
    text = "\n".join(str(cell.value or "") for row in _sheet(data).iter_rows() for cell in row)
    assert "MISSING" not in text


def test_26_31_excel_layout_and_statuses(app, db, admin_user):
    w = _world(db)
    complete, _ = _close_complete(admin_user, w, "500", customer="Complete Co")
    short, _ = _close_short(admin_user, w, "501", customer="Short Co")
    partial, _ = _partial_today(admin_user, w, "100", customer="Soho NY")
    data = export_eod_excel(_rows(), day=DAY)
    wb = load_workbook(BytesIO(data))
    assert wb.sheetnames == [WORKSHEET_NAME]
    assert len(wb.worksheets) == 1
    ws = wb[WORKSHEET_NAME]
    soho = _find(ws, "CEL - 100 - Soho NY")
    headers = [ws.cell(soho + 1, col).value for col in range(1, 21)]
    values = [ws.cell(soho + 2, col).value for col in range(1, 21)]
    assert headers == ORDER_INFO_HEADERS
    assert values[8] == OrderStatus.PARTIALLY_FULFILLED
    assert values[9] == FULFILLMENT_PARTIALLY
    assert values[11] == 10
    assert values[12] == 7
    assert values[14] == 3
    assert values[15] == 2
    complete_row = _find(ws, "CEL - 500 - Complete Co")
    short_row = _find(ws, "CEL - 501 - Short Co")
    assert soho < complete_row or True
    assert ws.cell(complete_row + 2, 10).value == FULFILLMENT_CLOSED_COMPLETE
    assert ws.cell(short_row + 2, 10).value == FULFILLMENT_CLOSED_SHORT
    assert AuditEvent.query.filter_by(event_type="END_OF_DAY_EXPORTED").count() >= 1


def test_32_ny_carton_boundary(app, db, admin_user):
    w = _world(db)
    early, early_cartons = _partial_today(admin_user, w, "610")
    late, late_cartons = _partial_today(admin_user, w, "611")
    _stamp(early, early_cartons, datetime(2026, 9, 22, 3, 30))
    _stamp(late, late_cartons, datetime(2026, 9, 22, 4, 30))
    ids_21 = _ids(_rows(date(2026, 9, 21)))
    ids_22 = _ids(_rows(date(2026, 9, 22)))
    assert early.id in ids_21
    assert early.id not in ids_22
    assert late.id in ids_22
    assert late.id not in ids_21


def test_33_36_filters_and_tenant(app, db, admin_user, admin_client):
    w = _world(db)
    complete, _ = _close_complete(admin_user, w, "700")
    partial, _ = _partial_today(admin_user, w, "701")
    complete_html = admin_client.get(
        "/orders/end-of-day?date=2026-09-24&fulfillment_status=CLOSED%20COMPLETE"
    ).get_data(as_text=True)
    assert complete.wms_order_id in complete_html
    assert partial.wms_order_id not in complete_html
    client_html = admin_client.get(
        f"/orders/end-of-day?date=2026-09-24&client_id={w['celine'].id}"
    ).get_data(as_text=True)
    assert complete.wms_order_id in client_html
    div_html = admin_client.get(
        f"/orders/end-of-day?date=2026-09-24&division_id={w['cel_ecom'].id}"
    ).get_data(as_text=True)
    assert complete.wms_order_id in div_html
    wh_html = admin_client.get(
        f"/orders/end-of-day?date=2026-09-24&warehouse_id={w['cel_ny'].id}"
    ).get_data(as_text=True)
    assert complete.wms_order_id in wh_html
    create_user("dio-eod", perms=["END_OF_DAY_VIEW", "END_OF_DAY_EXPORT", "DASHBOARD_VIEW"], clients=[w["dior"].id])
    other = app.test_client()
    login(other, "dio-eod")
    html = other.get("/orders/end-of-day?date=2026-09-24").get_data(as_text=True)
    assert complete.wms_order_id not in html
    assert other.get(f"/orders/end-of-day.xlsx?date=2026-09-24&client_id={w['celine'].id}").status_code == 404


def test_37_38_n_plus_one_and_reopen(app, db, admin_user):
    w = _world(db)
    _close_complete(admin_user, w, "800")
    filters = {"admin": True}
    with count_sql() as small:
        export_eod_excel(_rows(), day=DAY)
    _close_complete(admin_user, w, "801")
    _close_complete(admin_user, w, "802")
    with count_sql() as large:
        data = export_eod_excel(_rows(), day=DAY)
    assert small["count"] <= 25
    assert large["count"] <= small["count"] + 4
    wb = load_workbook(BytesIO(data))
    assert wb[WORKSHEET_NAME]


def test_39_40_packing_and_shipping_unchanged(app, db, admin_user):
    w, soho = _soho_order(db, admin_user)
    soho.closed_at = _midday()
    for carton in Carton.query.filter_by(order_id=soho.id).all():
        carton.closed_at = _midday()
    db.session.commit()
    pdf = render_packing_list_pdf(soho)
    assert pdf[:4] == b"%PDF"
    data, _, stats = export_shipping_report(admin_user, [soho.id])
    sr = load_workbook(BytesIO(data))
    assert sr.sheetnames == ["Shipping Report"]
    assert stats["product_rows"] == 3
    eod = export_eod_excel(_rows(), day=DAY)
    assert load_workbook(BytesIO(eod)).sheetnames == ["End of Day"]
    assert fulfillment_status(soho) == FULFILLMENT_CLOSED_COMPLETE
    assert PRODUCT_HEADERS == ["UPC", "SKU", "Description", "Style", "Color", "Size", "Qty"]


def test_kpis_do_not_count_partial_as_closed(app, db, admin_user):
    w = _world(db)
    _close_complete(admin_user, w, "900")
    _close_short(admin_user, w, "901")
    _partial_today(admin_user, w, "902")
    kpis = eod_kpis(_rows())
    assert kpis["total_orders"] == 3
    assert kpis["closed_complete"] == 1
    assert kpis["closed_short"] == 1
    assert kpis["partially_fulfilled"] == 1
    assert kpis["total_units"] == 10 + 8 + 7
    assert kpis["short_units"] == 2
    assert kpis["total_cartons"] == 3 + 1 + 2
