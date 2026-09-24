"""Shipping tracking, packing list, and End of Day operational date tests."""

from datetime import date, datetime, timezone
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.constants import (
    OrderStatus,
    PickTicketStatus,
    ShippingLabelStatus,
    ShippingStatus,
)
from app.extensions import db as _db
from app.models import AuditEvent, Carton, Document, InventoryUnit, Order, PickTicket
from app.services.allocation import allocate_order
from app.services.documents import get_store
from app.services.end_of_day import (
    build_eod_rows,
    closed_orders_query,
    eod_kpis,
    export_eod_excel,
    ny_day_utc_bounds,
    parse_eod_date,
)
from app.services.inventory_ledger import create_available_unit
from app.services.packing_list import existing_packing_list_document, persist_packing_list_pdf, render_packing_list_pdf
from app.services.pick_tickets import create_pick_ticket
from app.services.processing import (
    ProcessingError,
    acquire_lock,
    close_order,
    ensure_open_carton,
    find_ticket,
    request_close,
    scan_upc,
    set_dimensions,
    set_weight,
)
from app.services.shipping import ShippingError, confirm_shipping, save_carton_tracking
from tests.conftest import create_user, form_data, login
from tests.pdf_support import extract_pdf_text, pdf_page_count
from tests.test_v2_phase03_orders import _line
from tests.test_v2_phase04_allocation import _import_order
from tests.test_v2_phase05_pick_tickets import _ready_order
from tests.test_v2_phase06_processing import _processing_order
from tests.test_v2_phase07_reports import _closed


def _ready_order_in(db, admin_user, world, order_number, upc="UPC-A", qty=1, location="B-12"):
    """Create another allocated order without recreating unique clients."""
    for _ in range(qty):
        create_available_unit(
            client_id=world["celine"].id,
            warehouse_id=world["cel_ny"].id,
            upc=upc,
            location=location,
            sku="SKU-A",
        )
    db.session.commit()
    order = _import_order(
        admin_user,
        world["celine"],
        world["cel_ecom"],
        [_line(order=order_number, upc=upc, qty=qty)],
    )
    allocate_order(order, user=admin_user)
    return Order.query.get(order.id)


def _pack_close_cartons(admin_user, order, upc="UPC-A", carton_splits=None):
    acquire_lock(order, admin_user)
    remaining = sum(line.qty_ordered for line in order.lines)
    splits = carton_splits or [remaining]
    for qty in splits:
        carton = ensure_open_carton(order, admin_user)
        set_dimensions(carton, 12, 12, 12)
        _db.session.commit()
        for _ in range(qty):
            scan_upc(order, carton, upc, admin_user)
            remaining -= 1
        request_close(carton, admin_user)
        set_weight(carton, 4.3, user=admin_user)
    close_order(Order.query.get(order.id), admin_user)
    return Order.query.get(order.id)


def test_shipping_a_closed_order_pending_tracking(app, db, admin_user):
    _, order = _closed(db, admin_user)
    assert order.status == OrderStatus.CLOSED
    assert order.shipping_status == ShippingStatus.PENDING_TRACKING
    packing = Document.query.filter_by(order_id=order.id, type="PACKING_LIST").one()
    assert packing.storage_key
    assert get_store().open(packing.storage_key)[:4] == b"%PDF"


def test_shipping_b_open_order_cannot_confirm(app, db, admin_user):
    _, order, _, carton = _processing_order(db, admin_user, qty=1)
    with pytest.raises(ShippingError, match="CLOSED"):
        confirm_shipping(order, admin_user)
    assert order.status != OrderStatus.CLOSED
    assert order.shipping_status == ShippingStatus.NOT_READY


def test_shipping_c_d_e_f_multi_carton_tracking(app, db, admin_user):
    w, order = _ready_order(db, admin_user, order_number="5101", qty=2)
    ticket = create_pick_ticket(order)
    closed = _pack_close_cartons(admin_user, Order.query.get(order.id), carton_splits=[1, 1])
    cartons = Carton.query.filter_by(order_id=closed.id).order_by(Carton.id).all()
    assert len(cartons) == 2
    save_carton_tracking(closed, cartons[0], tracking_number="1ZAAA1111111111111", carrier="UPS", user=admin_user)
    db.session.commit()
    with pytest.raises(ShippingError, match="Every carton must have"):
        confirm_shipping(Order.query.get(closed.id), admin_user)
    save_carton_tracking(
        Order.query.get(closed.id),
        cartons[1],
        tracking_number="1ZBBB2222222222222",
        carrier="FEDEX",
        user=admin_user,
    )
    db.session.commit()
    confirm_shipping(Order.query.get(closed.id), admin_user)
    closed = Order.query.get(closed.id)
    cartons = Carton.query.filter_by(order_id=closed.id).order_by(Carton.id).all()
    assert cartons[0].tracking_number == "1ZAAA1111111111111"
    assert cartons[1].tracking_number == "1ZBBB2222222222222"
    assert cartons[0].tracking_carrier == "UPS"
    assert cartons[1].tracking_carrier == "FEDEX"
    assert {c.shipping_label_status for c in cartons} == {ShippingLabelStatus.VALIDATED}
    assert closed.shipping_status == ShippingStatus.TRACKING_COMPLETE
    assert closed.status == OrderStatus.CLOSED
    assert ticket.pick_ticket_number
    assert PickTicket.query.get(ticket.id).status == PickTicketStatus.CLOSED
    assert AuditEvent.query.filter_by(event_type="ORDER_SHIPPING_CONFIRMED").count() >= 1


def test_shipping_g_i_correction_audited_does_not_reopen(app, db, admin_user):
    _, order = _closed(db, admin_user)
    carton = Carton.query.filter_by(order_id=order.id).one()
    save_carton_tracking(order, carton, tracking_number="1ZOLD0000000000000", carrier="UPS", user=admin_user)
    db.session.commit()
    confirm_shipping(Order.query.get(order.id), admin_user)
    save_carton_tracking(
        Order.query.get(order.id),
        Carton.query.get(carton.id),
        tracking_number="1ZNEW0000000000000",
        carrier="UPS",
        user=admin_user,
        reason="typo",
    )
    db.session.commit()
    order = Order.query.get(order.id)
    carton = Carton.query.get(carton.id)
    assert carton.tracking_number == "1ZNEW0000000000000"
    assert order.status == OrderStatus.CLOSED
    assert PickTicket.query.filter_by(order_id=order.id).one().status == PickTicketStatus.CLOSED
    detail = AuditEvent.query.filter_by(event_type="CARTON_TRACKING_UPDATED").order_by(AuditEvent.id.desc()).first()
    assert detail is not None
    assert "1ZOLD0000000000000" in detail.detail
    assert "1ZNEW0000000000000" in detail.detail


def test_shipping_h_closed_ticket_still_blocked(app, db, admin_user):
    _, order = _closed(db, admin_user)
    ticket = PickTicket.query.filter_by(order_id=order.id).one()
    carton = Carton.query.filter_by(order_id=order.id).one()
    save_carton_tracking(order, carton, tracking_number="1ZAAA1111111111111", carrier="UPS", user=admin_user)
    db.session.commit()
    confirm_shipping(Order.query.get(order.id), admin_user)
    with pytest.raises(ProcessingError, match="closed because the associated order is complete"):
        find_ticket(ticket.pick_ticket_number)
    with pytest.raises(ProcessingError, match="closed because the associated order is complete"):
        acquire_lock(Order.query.get(order.id), admin_user)


def test_shipping_j_k_l_tenant_csrf_unauthorized(app, db, admin_user, admin_client):
    w, order = _closed(db, admin_user)
    carton = Carton.query.filter_by(order_id=order.id).one()
    create_user(
        "dio",
        perms=["SHIPPING_VIEW", "SHIPPING_EDIT", "SHIPPING_CONFIRM", "DASHBOARD_VIEW"],
        clients=[w["dior"].id],
    )
    other = app.test_client()
    login(other, "dio")
    assert other.get(f"/orders/shipping/{order.id}").status_code == 404
    assert other.post(
        f"/orders/shipping/{order.id}/confirm",
        data=form_data(other),
    ).status_code == 404

    viewer = create_user(
        "viewonly",
        perms=["SHIPPING_VIEW", "DASHBOARD_VIEW", "ORDERS_VIEW"],
        clients=[w["celine"].id],
    )
    limited = app.test_client()
    login(limited, "viewonly")
    denied = limited.post(
        f"/orders/shipping/{order.id}/confirm",
        data=form_data(limited),
    )
    assert denied.status_code == 403

    forged = admin_client.post(
        f"/orders/shipping/{order.id}/confirm",
        data={"choice": "yes"},
    )
    assert forged.status_code in {400, 403}
    assert Order.query.get(order.id).shipping_status == ShippingStatus.PENDING_TRACKING

    save_carton_tracking(order, carton, tracking_number="1ZAAA1111111111111", carrier="UPS", user=admin_user)
    db.session.commit()
    ok = admin_client.post(
        f"/orders/shipping/{order.id}/confirm",
        data=form_data(admin_client),
        follow_redirects=True,
    )
    assert ok.status_code == 200
    assert Order.query.get(order.id).shipping_status == ShippingStatus.TRACKING_COMPLETE


def test_packing_a_through_h(app, db, admin_user):
    w, order = _ready_order(db, admin_user, order_number="5201", qty=2, location="L-02")
    for unit in InventoryUnit.query.filter_by(allocated_order_id=order.id):
        unit.description = "Navy Coat"
        unit.style = "COAT"
        unit.color = "NAVY"
        unit.size = "M"
        unit.sku = "SKU-A"
    db.session.commit()
    create_pick_ticket(order)
    closed = _pack_close_cartons(admin_user, Order.query.get(order.id), carton_splits=[1, 1])
    packing = existing_packing_list_document(closed)
    assert packing is not None
    data = get_store().open(packing.storage_key)
    assert pdf_page_count(data) == 2
    text = extract_pdf_text(data)
    for needle in (
        "WMS SYSTEM",
        "PACKING LIST",
        closed.wms_order_id,
        closed.client_order_number,
        "Navy Coat",
        "COAT",
        "NAVY",
        "SKU",
        "UPC",
        "Description",
        "Style",
        "Color",
        "Size",
        "Qty in Carton",
    ):
        assert needle.lower() in text.lower()
    cartons = Carton.query.filter_by(order_id=closed.id).order_by(Carton.id).all()
    assert cartons[0].carton_number in text
    assert cartons[1].carton_number in text
    single = render_packing_list_pdf(closed)
    assert pdf_page_count(single) == 2
    key_before = packing.storage_key
    save_carton_tracking(closed, cartons[0], tracking_number="1ZHIST111111111111", carrier="UPS", user=admin_user)
    db.session.commit()
    reused = persist_packing_list_pdf(Order.query.get(closed.id))
    assert reused.storage_key == key_before
    assert reused.id == packing.id


def test_eod_ny_day_bounds_use_local_midnight():
    start, end = ny_day_utc_bounds(date(2026, 9, 21))
    assert start.tzinfo is None
    assert end.tzinfo is None
    # 2026-09-21 is EDT (UTC-4): local midnight is 04:00 UTC.
    assert start == datetime(2026, 9, 21, 4, 0)
    assert end == datetime(2026, 9, 22, 4, 0)


def test_eod_a_through_i(app, db, admin_user, admin_client):
    w, order = _closed(db, admin_user)
    ny_day = parse_eod_date(None)
    start, end = ny_day_utc_bounds(ny_day)
    assert start.tzinfo is None
    order.closed_at = start + (end - start) / 2
    db.session.commit()

    # UTC instants around NY midnight: 03:30 UTC is still Sep 21 EDT; 04:30 UTC is Sep 22.
    early_utc = datetime(2026, 9, 22, 3, 30, tzinfo=timezone.utc).replace(tzinfo=None)
    late_utc = datetime(2026, 9, 22, 4, 30, tzinfo=timezone.utc).replace(tzinfo=None)
    order2 = _ready_order_in(db, admin_user, w, "5302", qty=1, location="B-22")
    create_pick_ticket(order2)
    closed2 = _pack_close_cartons(admin_user, Order.query.get(order2.id), carton_splits=[1])
    closed2.closed_at = early_utc
    db.session.commit()
    order3 = _ready_order_in(db, admin_user, w, "5303", qty=1, location="B-23")
    create_pick_ticket(order3)
    closed3 = _pack_close_cartons(admin_user, Order.query.get(order3.id), carton_splits=[1])
    closed3.closed_at = late_utc
    db.session.commit()

    sep21 = date(2026, 9, 21)
    sep22 = date(2026, 9, 22)
    rows_21 = build_eod_rows(closed_orders_query(day=sep21, admin=True).all())
    rows_22 = build_eod_rows(closed_orders_query(day=sep22, admin=True).all())
    ids_21 = {r["order"].id for r in rows_21}
    ids_22 = {r["order"].id for r in rows_22}
    assert closed2.id in ids_21
    assert closed2.id not in ids_22
    assert closed3.id in ids_22
    assert closed3.id not in ids_21

    today_html = admin_client.get("/orders/end-of-day").get_data(as_text=True)
    assert "America/New_York" in today_html
    assert "Total Orders Processed" in today_html
    assert "Fulfillment Status" in today_html

    rows = build_eod_rows(closed_orders_query(day=parse_eod_date(None), admin=True).all())
    assert all(r["order"].status == OrderStatus.CLOSED for r in rows)
    codes = [(r["order"].client.client_code, r["order"].division.code, r["order"].closed_at) for r in rows]
    assert codes == sorted(codes)

    for row in rows:
        assert row["reconciled"] is True
        assert row["ordered"] == row["allocated"] == row["packed"] == row["shipped"] == row["carton_units"]
        assert row["carton_count"] == len(row["cartons"])
        assert row["total_weight"] == sum((c.weight or 0) for c in row["cartons"])

    create_user("dio", perms=["END_OF_DAY_VIEW", "DASHBOARD_VIEW"], clients=[w["dior"].id])
    other = app.test_client()
    login(other, "dio")
    html = other.get("/orders/end-of-day").get_data(as_text=True)
    assert order.wms_order_id not in html

    data = export_eod_excel(rows)
    book = load_workbook(filename=BytesIO(data))
    assert book.sheetnames == ["End of Day"]
    sheet = book["End of Day"]
    text = "\n".join(str(cell.value or "") for row in sheet.iter_rows() for cell in row)
    assert "WMS Order ID" in text
    assert "Fulfillment Status" in text
    assert "Shipping Status" in text
    assert "Carton " in text
    assert "Tracking:" in text
    assert AuditEvent.query.filter_by(event_type="END_OF_DAY_EXPORTED").count() >= 1
    kpis = eod_kpis(rows)
    assert kpis["total_orders"] == len(rows)
    assert kpis["closed_complete"] + kpis["closed_short"] + kpis["partially_fulfilled"] == len(rows)
