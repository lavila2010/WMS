"""Inventory Transactions: paginated screen and unlimited filtered Excel export."""

from datetime import datetime, timezone
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from app.constants import EOD_TIMEZONE, LedgerType, UnitStatus
from app.extensions import db
from app.models import Carton, InventoryTransaction, InventoryUnit, Order
from app.services.inventory_ledger import create_available_unit
from app.services.inventory_transactions import export_transactions_xlsx, list_transaction_page
from app.services.pick_tickets import create_pick_ticket
from app.services.processing import acquire_lock, close_order, ensure_open_carton, request_close, scan_upc, set_dimensions, set_weight
from app.services.shipping import save_carton_tracking
from tests.conftest import create_user, login
from tests.sql_probe import count_sql
from tests.test_v2_phase03_orders import _line
from tests.test_v2_phase04_allocation import _import_order, _world
from tests.test_v2_phase05_pick_tickets import _ready_order
from tests.test_v2_shipping_eod_packing import _pack_close_cartons


TZ = ZoneInfo(EOD_TIMEZONE)


def _closed_ship(db, admin_user, world=None, number="8800", customer="Ledger Co"):
    if world is None:
        w, order = _ready_order(db, admin_user, order_number=number, qty=1)
        create_pick_ticket(order)
        closed = _pack_close_cartons(admin_user, Order.query.get(order.id), carton_splits=[1])
        carton = Carton.query.filter_by(order_id=closed.id).one()
        save_carton_tracking(closed, carton, tracking_number="1ZLEDGER001", carrier="UPS", user=admin_user)
        db.session.commit()
        return w, Order.query.get(closed.id), Carton.query.get(carton.id)
    create_available_unit(
        client_id=world["celine"].id,
        warehouse_id=world["cel_ny"].id,
        upc="UPC-LEDGER",
        location="A-01",
        sku="SKU-L",
        description="Ledger Bag",
        style="ST-L",
        color="Black",
        size="OS",
    )
    db.session.commit()
    order = _import_order(
        admin_user,
        world["celine"],
        world["cel_ecom"],
        [_line(order=number, upc="UPC-LEDGER", qty=1, customer=customer)],
    )
    from app.services.allocation import allocate_order

    allocate_order(order, user=admin_user)
    create_pick_ticket(Order.query.get(order.id))
    closed = _pack_close_cartons(admin_user, Order.query.get(order.id), upc="UPC-LEDGER", carton_splits=[1])
    carton = Carton.query.filter_by(order_id=closed.id).one()
    save_carton_tracking(closed, carton, tracking_number="1ZLEDGER001", carrier="UPS", user=admin_user)
    db.session.commit()
    return world, Order.query.get(closed.id), Carton.query.get(carton.id)


def _bulk_ship_rows(client, warehouse, count, *, upc, created_at, location="SHIP-LOC"):
    units = [
        InventoryUnit(
            client_id=client.id,
            warehouse_id=warehouse.id,
            upc=upc,
            sku="SKU-BULK",
            description="Bulk Ship",
            style="ST-B",
            color="Navy",
            size="M",
            location=location,
            status=UnitStatus.SHIPPED,
        )
        for _ in range(count)
    ]
    db.session.add_all(units)
    db.session.flush()
    db.session.add_all(
        [
            InventoryTransaction(
                client_id=client.id,
                warehouse_id=warehouse.id,
                inventory_unit_id=unit.id,
                upc=upc,
                location=location,
                transaction_type=LedgerType.SHIP,
                from_status=UnitStatus.PACKED,
                to_status=UnitStatus.SHIPPED,
                created_at=created_at,
                reference="bulk-ship",
            )
            for unit in units
        ]
    )
    db.session.commit()


def test_01_page_loads(app, db, admin_user, admin_client):
    resp = admin_client.get("/inventory/transactions")
    html = resp.get_data(as_text=True)
    assert resp.status_code == 200
    assert "Movement Ledger" in html
    assert "From Date" in html
    assert "Transaction Type" in html
    assert "Export Excel" in html
    assert 'option value="100" selected' in html
    assert "250" in html and "500" in html
    for kind in LedgerType.ALL:
        assert kind in html


def test_02_03_04_client_warehouse_filters(app, db, admin_user, admin_client):
    w, order, carton = _closed_ship(db, admin_user)
    create_available_unit(client_id=w["dior"].id, warehouse_id=w["dio_ny"].id, upc="UPC-DIO", location="D-01")
    db.session.commit()
    cel_html = admin_client.get(f"/inventory/transactions?client_id={w['celine'].id}").get_data(as_text=True)
    assert order.wms_order_id in cel_html
    assert "UPC-DIO" not in cel_html
    wh_html = admin_client.get(
        f"/inventory/transactions?client_id={w['celine'].id}&warehouse_id={w['cel_ny'].id}"
    ).get_data(as_text=True)
    assert "UPC-A" in wh_html or order.wms_order_id in wh_html
    crossed = admin_client.get(
        f"/inventory/transactions?client_id={w['celine'].id}&warehouse_id={w['dio_ny'].id}"
    )
    assert crossed.status_code == 200
    assert "UPC-DIO" not in crossed.get_data(as_text=True)


def test_05_06_upc_and_ship_filter(app, db, admin_user, admin_client):
    w, order, carton = _closed_ship(db, admin_user)
    html = admin_client.get("/inventory/transactions?upc=UPC-A&type=SHIP").get_data(as_text=True)
    assert "SHIP" in html
    assert order.wms_order_id in html
    data, _, _ = export_transactions_xlsx(admin_user, {"upc": "UPC-A", "type": "SHIP"})
    types = {row[11] for row in load_workbook(BytesIO(data)).active.iter_rows(min_row=2, values_only=True)}
    assert types == {"SHIP"}


def test_07_08_09_ny_date_filters(app, db, admin_user, admin_client):
    w = _world(db)
    in_local = datetime(2026, 9, 22, 23, 30, tzinfo=TZ).astimezone(timezone.utc).replace(tzinfo=None)
    out_local = datetime(2026, 9, 22, 0, 30, tzinfo=timezone.utc).replace(tzinfo=None)
    _bulk_ship_rows(w["celine"], w["cel_ny"], 1, upc="UPC-IN-DAY", created_at=in_local, location="NY-IN")
    _bulk_ship_rows(w["celine"], w["cel_ny"], 1, upc="UPC-OUT-DAY", created_at=out_local, location="NY-OUT")
    html = admin_client.get(
        f"/inventory/transactions?client_id={w['celine'].id}&type=SHIP&date_from=2026-09-22&date_to=2026-09-22"
    ).get_data(as_text=True)
    assert "UPC-IN-DAY" in html
    assert "UPC-OUT-DAY" not in html
    data, _, count = export_transactions_xlsx(
        admin_user,
        {
            "client_id": w["celine"].id,
            "warehouse_id": w["cel_ny"].id,
            "type": "SHIP",
            "date_from": "2026-09-22",
            "date_to": "2026-09-22",
        },
    )
    ws = load_workbook(BytesIO(data)).active
    upcs = [row[4].value for row in ws.iter_rows(min_row=2, max_col=5)]
    assert "UPC-IN-DAY" in upcs
    assert "UPC-OUT-DAY" not in upcs
    assert count == 1


def test_10_13_screen_paginated_not_unlimited(app, db, admin_user, admin_client):
    w = _world(db)
    now = datetime.utcnow()
    _bulk_ship_rows(w["celine"], w["cel_ny"], 120, upc="UPC-PAGE", created_at=now)
    html = admin_client.get(
        f"/inventory/transactions?client_id={w['celine'].id}&upc=UPC-PAGE&type=SHIP&per_page=100"
    ).get_data(as_text=True)
    assert "Movement Ledger (120)" in html
    assert html.count("UPC-PAGE") <= 120
    assert "page 1 of 2" in html
    page2 = admin_client.get(
        f"/inventory/transactions?client_id={w['celine'].id}&upc=UPC-PAGE&type=SHIP&per_page=100&page=2"
    ).get_data(as_text=True)
    assert "page 2 of 2" in page2


def test_12_14_18_export_all_over_500(app, db, admin_user, admin_client):
    w = _world(db)
    now = datetime.utcnow()
    _bulk_ship_rows(w["celine"], w["cel_ny"], 520, upc="UPC-MANY", created_at=now)
    page = list_transaction_page(
        admin_user,
        {"client_id": w["celine"].id, "upc": "UPC-MANY", "type": "SHIP"},
        page=1,
        per_page=100,
    )
    assert page["total"] == 520
    assert len(page["rows"]) == 100
    resp = admin_client.get(
        f"/inventory/transactions.xlsx?client_id={w['celine'].id}&warehouse_id={w['cel_ny'].id}&upc=UPC-MANY&type=SHIP"
    )
    assert resp.status_code == 200
    assert "Inventory_Transactions_" in (resp.headers.get("Content-Disposition") or "")
    wb = load_workbook(BytesIO(resp.data))
    assert wb.sheetnames == ["Transactions"]
    assert wb.active.max_row == 521
    types = {row[11] for row in wb.active.iter_rows(min_row=2, values_only=True)}
    assert types == {"SHIP"}


def test_15_16_17_export_matches_filters(app, db, admin_user):
    w, order, carton = _closed_ship(db, admin_user)
    other = datetime(2020, 1, 1, 12, 0, 0)
    _bulk_ship_rows(w["celine"], w["cel_nj"], 3, upc="UPC-NJ", created_at=other, location="NJ-01")
    data, _, count = export_transactions_xlsx(
        admin_user,
        {
            "client_id": w["celine"].id,
            "warehouse_id": w["cel_ny"].id,
            "type": "SHIP",
            "date_from": datetime.now(TZ).strftime("%Y-%m-%d"),
            "date_to": datetime.now(TZ).strftime("%Y-%m-%d"),
        },
    )
    ws = load_workbook(BytesIO(data)).active
    warehouses = {row[2] for row in ws.iter_rows(min_row=2, values_only=True)}
    clients = {row[1] for row in ws.iter_rows(min_row=2, values_only=True)}
    assert "01-CEL-NJ" not in warehouses
    assert clients <= {"01-CEL"}
    assert count >= 1


def test_19_24_ship_context(app, db, admin_user, admin_client):
    w, order, carton = _closed_ship(db, admin_user)
    html = admin_client.get(f"/inventory/transactions?type=SHIP&upc=UPC-A").get_data(as_text=True)
    assert "B-02" in html or "A-01" in html
    assert order.wms_order_id in html
    assert order.client_order_number in html
    assert "Ada" in html or (order.customer or "") in html
    assert carton.carton_number in html
    assert "1ZLEDGER001" in html
    data, _, _ = export_transactions_xlsx(admin_user, {"type": "SHIP", "upc": "UPC-A"})
    ws = load_workbook(BytesIO(data)).active
    headers = [cell.value for cell in ws[1]]
    assert headers == [
        "Timestamp",
        "Client",
        "Warehouse",
        "Unit ID",
        "UPC",
        "SKU",
        "Description",
        "Style",
        "Color",
        "Size",
        "Location",
        "Transaction",
        "From Status",
        "To Status",
        "WMS Order ID",
        "Client Order Number",
        "Customer",
        "Pick Ticket",
        "Carton Number",
        "Tracking Number",
        "Reference",
    ]
    values = [cell.value for cell in next(ws.iter_rows(min_row=2, max_row=2))]
    assert values[14] == order.wms_order_id
    assert values[15] == order.client_order_number
    assert values[16] == order.customer
    assert values[18] == carton.carton_number
    assert values[19] == "1ZLEDGER001"
    assert values[10]


def test_25_no_cross_client_export(app, db, admin_user):
    w, order, _carton = _closed_ship(db, admin_user)
    create_available_unit(client_id=w["dior"].id, warehouse_id=w["dio_ny"].id, upc="UPC-DIO", location="D-01")
    db.session.commit()
    user = create_user(
        "cel-txn",
        perms=["INVENTORY_VIEW", "INVENTORY_EXPORT"],
        clients=[w["celine"].id],
    )
    db.session.refresh(user)
    from app.services.inventory_transactions import TransactionAccessError
    import pytest

    with pytest.raises(TransactionAccessError):
        export_transactions_xlsx(user, {"client_id": w["dior"].id})
    client = app.test_client()
    login(client, "cel-txn")
    assert client.get(f"/inventory/transactions.xlsx?client_id={w['dior'].id}").status_code == 404
    ok = client.get(f"/inventory/transactions.xlsx?client_id={w['celine'].id}&type=SHIP")
    assert ok.status_code == 200
    text = " ".join(
        str(cell.value or "")
        for row in load_workbook(BytesIO(ok.data)).active.iter_rows()
        for cell in row
    )
    assert "UPC-DIO" not in text


def test_26_no_n_plus_one(app, db, admin_user):
    w = _world(db)
    now = datetime.utcnow()
    _bulk_ship_rows(w["celine"], w["cel_ny"], 40, upc="UPC-N1", created_at=now)
    filters = {"client_id": w["celine"].id, "upc": "UPC-N1", "type": "SHIP"}
    with count_sql() as small:
        export_transactions_xlsx(admin_user, filters)
    _bulk_ship_rows(w["celine"], w["cel_ny"], 80, upc="UPC-N1", created_at=now)
    with count_sql() as large:
        export_transactions_xlsx(admin_user, filters)
    assert small["count"] <= 25
    assert large["count"] <= small["count"] + 4


def test_27_excel_reopens(app, db, admin_user):
    w = _world(db)
    _bulk_ship_rows(w["celine"], w["cel_ny"], 3, upc="UPC-OPEN", created_at=datetime.utcnow())
    data, filename, count = export_transactions_xlsx(admin_user, {"upc": "UPC-OPEN"})
    wb = load_workbook(BytesIO(data))
    assert wb["Transactions"]
    assert filename.startswith("Inventory_Transactions_")
    assert count == 3
