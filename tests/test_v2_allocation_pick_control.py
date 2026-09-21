"""Bulk allocation, partial waves, daily exceptions, substitution, issues, bulk PDF."""

from __future__ import annotations

from datetime import datetime, timedelta
from io import BytesIO
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook

from app.constants import (
    CLOSED_PICK_TICKET_UPDATE_MESSAGE,
    REPLACEMENT_UNAVAILABLE_MESSAGE,
    AllocationStatus,
    InventoryIssueStatus,
    LedgerType,
    OrderStatus,
    PickTicketStatus,
    UnitStatus,
)
from app.extensions import db
from app.models import (
    Allocation,
    AuditEvent,
    InventoryIssue,
    InventoryTransaction,
    InventoryUnit,
    Order,
    OrderLine,
    PickTicket,
    PickTicketPrintBatch,
)
from app.services.allocation import (
    AllocationError,
    allocate_order,
    allocate_orders,
    approve_partial_allocation,
)
from app.services.allocation_exceptions import export_daily_exceptions, list_daily_exceptions
from app.services.bulk_pick_pdf import publish_bulk_pdf, render_bulk_pdf
from app.services.end_of_day import ny_day_utc_bounds, operational_today
from app.services.fulfillment import order_quantities
from app.services.inventory_issues import confirm_missing, decommission_unit, resolve_found
from app.services.inventory_ledger import create_available_unit
from app.services.inventory_query import status_counts
from app.services.pick_ticket_update import PickTicketUpdateError, substitute_unit
from app.services.pick_tickets import PickTicketError, create_pick_ticket, render_pdf, ticket_lines
from app.services.processing import (
    acquire_lock,
    close_order,
    ensure_open_carton,
    find_ticket,
    request_close,
    scan_upc,
    set_dimensions,
    set_weight,
)
from tests.conftest import create_user, form_data, login
from tests.pdf_support import extract_pdf_text
from tests.test_v2_phase03_orders import _line
from tests.test_v2_phase04_allocation import _import_order, _stock, _world
from tests.test_v2_phase05_pick_tickets import _ready_order


def _order(admin_user, w, number, qty, upc="UPC-A", warehouse="NY"):
    return _import_order(
        admin_user,
        w["celine"],
        w["cel_ecom"],
        [_line(order=str(number), warehouse=warehouse, upc=upc, qty=qty)],
    )


def _pack_ticket(admin_user, order, upc="UPC-A", qty=None):
    acquire_lock(order, admin_user)
    carton = ensure_open_carton(order, admin_user)
    set_dimensions(carton, 12, 12, 12)
    db.session.commit()
    count = qty if qty is not None else order_quantities(order)["currently_allocated"]
    for _ in range(count):
        scan_upc(order, carton, upc, admin_user)
    request_close(carton, admin_user)
    set_weight(carton, 2.2, user=admin_user)
    close_order(Order.query.get(order.id), admin_user)
    return Order.query.get(order.id)


def test_ay_01_single_and_multi_and_select_visible(app, db, admin_user, admin_client):
    w = _world(db)
    _stock(w["celine"], w["cel_ny"], "UPC-A", 6)
    o1 = _order(admin_user, w, 4101, 2)
    o2 = _order(admin_user, w, 4102, 2)
    html = admin_client.get("/allocation/").get_data(as_text=True)
    assert "Select All Visible" in html
    assert "Allocate Selected" in html
    one = allocate_orders([o1.id], user=admin_user)
    assert one["selected"] == 1
    assert one["fully_allocated"] == 1
    two = allocate_orders([o2.id], user=admin_user)
    assert two["fully_allocated"] == 1
    resp = admin_client.post(
        "/allocation/bulk",
        data=form_data(admin_client, {"order_ids": [str(o1.id)]}),
        follow_redirects=True,
    )
    assert resp.status_code == 200


def test_ay_04_05_closed_cancelled_blocked(app, db, admin_user):
    w, order = _ready_order(db, admin_user, order_number="4103", qty=1)
    create_pick_ticket(order)
    _pack_ticket(admin_user, Order.query.get(order.id), qty=1)
    closed = Order.query.get(order.id)
    assert closed.status == OrderStatus.CLOSED
    with pytest.raises(AllocationError):
        allocate_order(closed, user=admin_user)
    cancelled = _order(admin_user, w, 4104, 1)
    cancelled.status = OrderStatus.CANCELLED
    db.session.commit()
    with pytest.raises(AllocationError):
        allocate_order(cancelled, user=admin_user)


def test_ay_06_07_independent_and_no_double_unit(app, db, admin_user):
    w = _world(db)
    _stock(w["celine"], w["cel_ny"], "UPC-A", 2)
    ok = _order(admin_user, w, 4105, 1)
    fail = _order(admin_user, w, 4106, 1)
    fail.status = OrderStatus.CANCELLED
    db.session.commit()
    summary = allocate_orders([ok.id, fail.id], user=admin_user)
    assert summary["fully_allocated"] == 1
    assert summary["failed"] == 1
    assert InventoryUnit.query.filter_by(status=UnitStatus.RESERVED).count() == 1
    unit_ids = [a.inventory_unit_id for a in Allocation.query.filter_by(status=AllocationStatus.ACTIVE)]
    assert len(unit_ids) == len(set(unit_ids))


def test_ay_08_09_tenant_and_permission(app, db, admin_user, client):
    w = _world(db)
    _stock(w["celine"], w["cel_ny"], "UPC-A", 1)
    order = _order(admin_user, w, 4107, 1)
    create_user("dio", perms=["ALLOCATION_VIEW"], clients=[w["dior"].id])
    login(client, "dio")
    assert client.get(f"/allocation/{order.id}").status_code == 404
    create_user("noperm", perms=["DASHBOARD_VIEW"], clients=[w["celine"].id])
    login(client, "noperm")
    assert client.post(
        "/allocation/bulk", data=form_data(client, {"order_ids": [str(order.id)]})
    ).status_code == 403


def test_az_10_19_partial_waves(app, db, admin_user):
    w = _world(db)
    _stock(w["celine"], w["cel_ny"], "UPC-A", 10)
    order = _order(admin_user, w, 4201, 10)
    # consume 3 so first allocate can only get 7
    for unit in InventoryUnit.query.filter_by(upc="UPC-A").limit(3):
        unit.status = UnitStatus.SHIPPED
    db.session.commit()
    result = allocate_order(Order.query.get(order.id), user=admin_user)
    order = Order.query.get(order.id)
    assert result["reserved"] == 7
    assert order.status == OrderStatus.PARTIALLY_ALLOCATED
    assert PickTicket.query.filter_by(order_id=order.id).count() == 0
    with pytest.raises(PickTicketError):
        create_pick_ticket(order)
    approve_partial_allocation(order, user=admin_user)
    ticket1 = create_pick_ticket(Order.query.get(order.id))
    assert ticket1.pick_ticket_number.endswith("-01")
    lines = ticket_lines(Order.query.get(order.id), ticket1)
    assert sum(line["qty"] for line in lines) == 7
    qty = order_quantities(Order.query.get(order.id))
    assert qty["remaining"] == 3
    _pack_ticket(admin_user, Order.query.get(order.id), qty=7)
    order = Order.query.get(order.id)
    assert ticket1.status == PickTicketStatus.CLOSED or PickTicket.query.get(ticket1.id).status == PickTicketStatus.CLOSED
    assert order.status == OrderStatus.PARTIALLY_FULFILLED
    assert order_quantities(order)["remaining"] == 3
    for unit in InventoryUnit.query.filter_by(status=UnitStatus.SHIPPED, upc="UPC-A").limit(3):
        if unit.allocated_order_id is None:
            unit.status = UnitStatus.AVAILABLE
    db.session.commit()
    allocate_order(Order.query.get(order.id), user=admin_user)
    ticket2 = create_pick_ticket(Order.query.get(order.id))
    assert ticket2.ticket_sequence == 2
    assert ticket2.pick_ticket_number.endswith("-02")
    t2_units = {row["upc"] for row in ticket_lines(Order.query.get(order.id), ticket2)}
    shipped_ids = {
        u.id for u in InventoryUnit.query.filter_by(allocated_order_id=order.id, status=UnitStatus.SHIPPED)
    }
    t2_alloc = Allocation.query.filter_by(pick_ticket_id=ticket2.id).all()
    assert not {a.inventory_unit_id for a in t2_alloc} & shipped_ids
    _pack_ticket(admin_user, Order.query.get(order.id), qty=3)
    final = Order.query.get(order.id)
    assert final.status == OrderStatus.CLOSED
    assert order_quantities(final)["shipped"] == 10
    assert AuditEvent.query.filter_by(event_type="PARTIAL_ALLOCATION_APPROVED").count() >= 1
    assert AuditEvent.query.filter_by(event_type="ORDER_FULLY_FULFILLED").count() >= 1


def test_ba_20_33_daily_exceptions(app, db, admin_user, admin_client):
    w = _world(db)
    _stock(w["celine"], w["cel_ny"], "UPC-A", 8)
    unalloc = _order(admin_user, w, 4301, 2)
    partial = _order(admin_user, w, 4302, 4)
    for unit in InventoryUnit.query.filter_by(upc="UPC-A").limit(6):
        unit.status = UnitStatus.SHIPPED
    db.session.commit()
    allocate_order(partial, user=admin_user)
    for unit in InventoryUnit.query.filter_by(upc="UPC-A", status=UnitStatus.SHIPPED):
        if unit.allocated_order_id is None:
            unit.status = UnitStatus.AVAILABLE
    db.session.commit()
    # force partial by adding another line remaining... already 4 allocated if stock enough
    full = _order(admin_user, w, 4303, 1)
    allocate_order(full, user=admin_user)
    closed = _order(admin_user, w, 4304, 1)
    allocate_order(closed, user=admin_user)
    create_pick_ticket(Order.query.get(closed.id))
    _pack_ticket(admin_user, Order.query.get(closed.id), qty=1)
    cancelled = _order(admin_user, w, 4305, 1)
    cancelled.status = OrderStatus.CANCELLED
    db.session.commit()
    payload = list_daily_exceptions(admin=True)
    ids = {row["order"].id for row in payload["rows"]}
    assert unalloc.id in ids
    assert closed.id not in ids
    assert cancelled.id not in ids
    assert Order.query.get(full.id).id not in ids or order_quantities(Order.query.get(full.id))["remaining"] == 0
    assert payload["kpis"]["outstanding_units"] == sum(r["remaining"] for r in payload["rows"])
    assert len(payload["rows"]) == len({r["order"].id for r in payload["rows"]})
    data = export_daily_exceptions(payload)
    book = load_workbook(BytesIO(data))
    assert "Orders" in book.sheetnames
    assert "Outstanding by UPC" in book.sheetnames
    html = admin_client.get("/allocation/exceptions").get_data(as_text=True)
    assert "Daily Allocation Exceptions" in html
    tz = ZoneInfo("America/New_York")
    start, _end = ny_day_utc_bounds(operational_today())
    yesterday = start - timedelta(hours=1)
    old = _order(admin_user, w, 4306, 1)
    old.created_at = yesterday
    db.session.commit()
    today_ids = {row["order"].id for row in list_daily_exceptions(admin=True)["rows"]}
    assert old.id not in today_ids


def test_bb_34_54_substitution(app, db, admin_user):
    w, order = _ready_order(db, admin_user, order_number="4401", qty=2, location="A-01")
    extra = create_available_unit(
        client_id=w["celine"].id, warehouse_id=w["cel_ny"].id, upc="UPC-A", location="B-99", sku="SKU-A"
    )
    db.session.commit()
    ticket = create_pick_ticket(order)
    reserved = InventoryUnit.query.filter_by(allocated_order_id=order.id, status=UnitStatus.RESERVED).first()
    before_alloc = order_quantities(order)["currently_allocated"]
    result = substitute_unit(ticket, reserved, extra, user=admin_user)
    assert result["revision"] == 2
    assert PickTicket.query.get(ticket.id).pick_ticket_number == ticket.pick_ticket_number
    assert InventoryUnit.query.get(reserved.id).status == UnitStatus.ISSUE_HOLD
    assert InventoryUnit.query.get(extra.id).status == UnitStatus.RESERVED
    assert order_quantities(Order.query.get(order.id))["currently_allocated"] == before_alloc
    assert InventoryIssue.query.count() == 1
    text = extract_pdf_text(render_pdf(PickTicket.query.get(ticket.id)))
    assert "B-99" in text
    assert "Rev 2" in text
    closed = _ready_order(db, admin_user, order_number="4402", qty=1)
    t2 = create_pick_ticket(closed[1])
    _pack_ticket(admin_user, closed[1], qty=1)
    t2 = PickTicket.query.get(t2.id)
    with pytest.raises(PickTicketUpdateError, match="cannot be updated"):
        substitute_unit(t2, InventoryUnit.query.first(), extra, user=admin_user)
    cancelled = _ready_order(db, admin_user, order_number="4403", qty=1)
    t3 = create_pick_ticket(cancelled[1])
    t3.status = PickTicketStatus.CANCELLED
    db.session.commit()
    with pytest.raises(PickTicketUpdateError):
        substitute_unit(t3, InventoryUnit.query.first(), extra, user=admin_user)


def test_bb_53_54_concurrent_replacement(app, db, admin_user):
    from sqlalchemy import text

    w, order = _ready_order(db, admin_user, order_number="4404", qty=1, location="A-01")
    repl = create_available_unit(
        client_id=w["celine"].id, warehouse_id=w["cel_ny"].id, upc="UPC-A", location="C-01"
    )
    db.session.commit()
    ticket = create_pick_ticket(order)
    original = InventoryUnit.query.filter_by(allocated_order_id=order.id, status=UnitStatus.RESERVED).one()
    with db.engine.connect() as held:
        trans = held.begin()
        held.execute(text("SELECT id FROM inventory_units WHERE id = :id FOR UPDATE"), {"id": repl.id})
        with pytest.raises(PickTicketUpdateError, match="no longer available"):
            substitute_unit(ticket, original, InventoryUnit.query.get(repl.id), user=admin_user)
        trans.rollback()
    db.session.rollback()
    original = InventoryUnit.query.get(original.id)
    repl = InventoryUnit.query.get(repl.id)
    assert original.status == UnitStatus.RESERVED
    assert repl.status == UnitStatus.AVAILABLE
    assert InventoryIssue.query.count() == 0
    substitute_unit(ticket, original, repl, user=admin_user)
    with pytest.raises(PickTicketUpdateError):
        substitute_unit(
            PickTicket.query.get(ticket.id),
            InventoryUnit.query.get(original.id),
            InventoryUnit.query.get(repl.id),
            user=admin_user,
        )
    assert InventoryUnit.query.filter_by(status=UnitStatus.RESERVED, upc="UPC-A").count() == 1
    assert InventoryIssue.query.count() == 1


def test_bc_55_62_issues_and_on_hand(app, db, admin_user, client):
    w, order = _ready_order(db, admin_user, order_number="4501", qty=1)
    repl = create_available_unit(
        client_id=w["celine"].id, warehouse_id=w["cel_ny"].id, upc="UPC-A", location="Z-01"
    )
    db.session.commit()
    ticket = create_pick_ticket(order)
    original = InventoryUnit.query.filter_by(allocated_order_id=order.id, status=UnitStatus.RESERVED).one()
    substitute_unit(ticket, original, repl, user=admin_user)
    issue = InventoryIssue.query.one()
    counts = status_counts(client_id=w["celine"].id)
    assert InventoryUnit.query.get(original.id).status == UnitStatus.ISSUE_HOLD
    assert counts["on_hand"] == status_counts(client_id=w["celine"].id)["available"] + status_counts(client_id=w["celine"].id)["reserved"] + status_counts(client_id=w["celine"].id)["packed"]
    with pytest.raises(Exception):
        resolve_found(issue, location="")
    resolve_found(InventoryIssue.query.get(issue.id), location="FOUND-1")
    assert InventoryUnit.query.get(original.id).status == UnitStatus.AVAILABLE
    assert InventoryUnit.query.get(original.id).location == "FOUND-1"
    assert InventoryTransaction.query.filter_by(transaction_type=LedgerType.ISSUE_RESOLVED).count() >= 1
    # second issue path
    hold = create_available_unit(
        client_id=w["celine"].id, warehouse_id=w["cel_ny"].id, upc="UPC-B", location="H-01"
    )
    db.session.commit()
    from app.services.inventory_ledger import transition_unit

    transition_unit(hold, to_status=UnitStatus.RESERVED, transaction_type=LedgerType.RESERVE, allocated_order_id=order.id)
    # force hold
    transition_unit(hold, to_status=UnitStatus.ISSUE_HOLD, transaction_type=LedgerType.ISSUE_HOLD)
    issue2 = InventoryIssue(
        client_id=w["celine"].id,
        warehouse_id=w["cel_ny"].id,
        inventory_unit_id=hold.id,
        upc="UPC-B",
        original_location="H-01",
        issue_type="PICK_UNIT_NOT_FOUND",
        status="OPEN",
    )
    db.session.add(issue2)
    db.session.commit()
    confirm_missing(issue2)
    assert InventoryUnit.query.get(hold.id).status == UnitStatus.MISSING
    decommission_unit(InventoryIssue.query.get(issue2.id))
    assert InventoryUnit.query.get(hold.id).status == UnitStatus.DECOMMISSIONED
    assert InventoryUnit.query.get(hold.id) is not None
    create_user("viewonly", perms=["INVENTORY_VIEW"], clients=[w["celine"].id])
    login(client, "viewonly")
    assert client.get("/inventory/issues").status_code == 403


def test_bd_63_77_bulk_pdf(app, db, admin_user):
    w = _world(db)
    _stock(w["celine"], w["cel_ny"], "UPC-A", 5)
    tickets = []
    for idx, number in enumerate((4601, 4602, 4603), start=1):
        order = _order(admin_user, w, number, 1)
        allocate_order(order, user=admin_user)
        tickets.append(create_pick_ticket(Order.query.get(order.id)))
    one = render_bulk_pdf(tickets[:1])
    assert one.startswith(b"%PDF")
    many = render_bulk_pdf(tickets, sort="ticket", direction="asc")
    assert many.startswith(b"%PDF")
    text = extract_pdf_text(many)
    assert "BULK PICK TICKET PACKAGE" in text
    for ticket in tickets:
        assert ticket.pick_ticket_number in text
    assert PickTicket.query.count() == 3
    before = {t.id: t.print_count for t in tickets}
    published = publish_bulk_pdf(tickets, sort="ticket", direction="asc")
    assert published["batch"].ticket_count == 3
    assert PickTicketPrintBatch.query.count() == 1
    for ticket in tickets:
        assert PickTicket.query.get(ticket.id).print_count == before[ticket.id] + 1
    closed_order = _order(admin_user, w, 4604, 1)
    allocate_order(closed_order, user=admin_user)
    closed_ticket = create_pick_ticket(Order.query.get(closed_order.id))
    _pack_ticket(admin_user, Order.query.get(closed_order.id), qty=1)
    closed_ticket = PickTicket.query.get(closed_ticket.id)
    assert closed_ticket.status == PickTicketStatus.CLOSED
    reprint = render_bulk_pdf([closed_ticket])
    assert b"%PDF" == reprint[:4]
    assert PickTicket.query.get(closed_ticket.id).status == PickTicketStatus.CLOSED


def test_bd_75_failed_batch_does_not_print(app, db, admin_user, monkeypatch):
    w, order = _ready_order(db, admin_user, order_number="4605", qty=1)
    ticket = create_pick_ticket(order)
    before = ticket.print_count

    def boom(_ticket):
        raise RuntimeError("render failed")

    monkeypatch.setattr("app.services.bulk_pick_pdf.render_pdf", boom)
    from app.services.bulk_pick_pdf import BulkPickPdfError

    with pytest.raises(BulkPickPdfError):
        render_bulk_pdf([ticket])
    assert PickTicket.query.get(ticket.id).print_count == before


def test_bd_78_hundred_ticket_batch(app, db, admin_user):
    import time

    w = _world(db)
    _stock(w["celine"], w["cel_ny"], "UPC-H", 100)
    tickets = []
    for i in range(100):
        order = _order(admin_user, w, 5000 + i, 1, upc="UPC-H")
        allocate_order(order, user=admin_user)
        tickets.append(create_pick_ticket(Order.query.get(order.id)))
    started = time.perf_counter()
    data = render_bulk_pdf(tickets, sort="ticket", direction="asc")
    elapsed = time.perf_counter() - started
    assert data.startswith(b"%PDF")
    assert elapsed < 60


def test_schema_backfill_does_not_attach_later_wave(app, db, admin_user):
    from app.schema import ensure_v2_schema

    w = _world(db)
    _stock(w["celine"], w["cel_ny"], "UPC-A", 10)
    order = _order(admin_user, w, 4801, 10)
    for unit in InventoryUnit.query.filter_by(upc="UPC-A").limit(3):
        unit.status = UnitStatus.SHIPPED
    db.session.commit()
    allocate_order(Order.query.get(order.id), user=admin_user)
    approve_partial_allocation(Order.query.get(order.id), user=admin_user)
    ticket1 = create_pick_ticket(Order.query.get(order.id))
    ticket1_id = ticket1.id
    order_id = order.id
    for unit in InventoryUnit.query.filter_by(status=UnitStatus.SHIPPED, upc="UPC-A"):
        if unit.allocated_order_id is None:
            unit.status = UnitStatus.AVAILABLE
    db.session.commit()
    allocate_order(Order.query.get(order.id), user=admin_user)
    db.session.commit()
    db.session.close()
    ensure_v2_schema()
    pending = Allocation.query.filter_by(
        order_id=order_id, status=AllocationStatus.ACTIVE, pick_ticket_id=None
    ).count()
    assert pending == 3
    assert Allocation.query.filter_by(pick_ticket_id=ticket1_id, status=AllocationStatus.ACTIVE).count() == 7
    ticket2 = create_pick_ticket(Order.query.get(order_id))
    assert ticket2.id != ticket1_id
    assert ticket2.ticket_sequence == 2
    assert Allocation.query.filter_by(pick_ticket_id=ticket2.id, status=AllocationStatus.ACTIVE).count() == 3


def test_be_end_to_end_wave(app, db, admin_user):
    w = _world(db)
    _stock(w["celine"], w["cel_ny"], "UPC-A", 10, location="A-01")
    db.session.commit()
    order = _order(admin_user, w, 4701, 10)
    for unit in InventoryUnit.query.filter_by(upc="UPC-A", location="A-01").limit(3):
        unit.status = UnitStatus.SHIPPED
    db.session.commit()
    allocate_order(order, user=admin_user)
    extra = create_available_unit(
        client_id=w["celine"].id, warehouse_id=w["cel_ny"].id, upc="UPC-A", location="R-01"
    )
    db.session.commit()
    order = Order.query.get(order.id)
    assert order.status == OrderStatus.PARTIALLY_ALLOCATED
    approve_partial_allocation(order, user=admin_user)
    t1 = create_pick_ticket(order)
    assert t1.pick_ticket_number.endswith("-01")
    original = InventoryUnit.query.filter_by(allocated_order_id=order.id, status=UnitStatus.RESERVED).first()
    substitute_unit(t1, original, extra, user=admin_user)
    assert order_quantities(Order.query.get(order.id))["currently_allocated"] == 7
    _pack_ticket(admin_user, Order.query.get(order.id), qty=7)
    order = Order.query.get(order.id)
    assert order.status == OrderStatus.PARTIALLY_FULFILLED
    assert order_quantities(order)["remaining"] == 3
    payload = list_daily_exceptions(admin=True)
    match = [row for row in payload["rows"] if row["order"].id == order.id]
    assert match and match[0]["remaining"] == 3
    for unit in InventoryUnit.query.filter_by(status=UnitStatus.SHIPPED, upc="UPC-A"):
        if unit.allocated_order_id is None:
            unit.status = UnitStatus.AVAILABLE
    db.session.commit()
    allocate_order(Order.query.get(order.id), user=admin_user)
    t2 = create_pick_ticket(Order.query.get(order.id))
    assert t2.pick_ticket_number.endswith("-02")
    render_bulk_pdf([t2, t1], sort="ticket", direction="asc")
    _pack_ticket(admin_user, Order.query.get(order.id), qty=3)
    final = Order.query.get(order.id)
    assert final.status == OrderStatus.CLOSED
    assert order_quantities(final)["shipped"] == 10
    assert PickTicket.query.filter_by(order_id=final.id).count() == 2
