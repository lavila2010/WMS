"""Default inventory views hide ON_HAND = 0 unless show_zero=1."""

from io import BytesIO

from openpyxl import load_workbook

from app.constants import LedgerType, UnitStatus
from app.models import InventoryTransaction, InventoryUnit
from app.services.inventory_ledger import create_available_unit, transition_unit
from app.services.inventory_query import aggregate_rows, search_units, status_counts
from app.services.masters import create_client, create_warehouse


def _ship(unit):
    transition_unit(unit, to_status=UnitStatus.RESERVED, transaction_type=LedgerType.RESERVE)
    transition_unit(unit, to_status=UnitStatus.PACKED, transaction_type=LedgerType.PACK)
    transition_unit(unit, to_status=UnitStatus.SHIPPED, transaction_type=LedgerType.SHIP)


def _world(db):
    celine = create_client("Celine", "CEL")
    db.session.commit()
    ny = create_warehouse(celine, "NY")
    nj = create_warehouse(celine, "NJ")
    db.session.commit()
    return celine, ny, nj


def _unit(celine, warehouse, upc, location, sku="SKU-A"):
    return create_available_unit(
        client_id=celine.id,
        warehouse_id=warehouse.id,
        upc=upc,
        location=location,
        sku=sku,
        description="Navy Coat",
    )


def test_a_b_c_available_reserved_packed_shown(app, db, admin_user):
    celine, ny, _ = _world(db)
    available = _unit(celine, ny, "UPC-AVAIL", "A-01")
    reserved = _unit(celine, ny, "UPC-RES", "R-01")
    packed = _unit(celine, ny, "UPC-PACK", "P-01")
    transition_unit(reserved, to_status=UnitStatus.RESERVED, transaction_type=LedgerType.RESERVE)
    transition_unit(packed, to_status=UnitStatus.RESERVED, transaction_type=LedgerType.RESERVE)
    transition_unit(packed, to_status=UnitStatus.PACKED, transaction_type=LedgerType.PACK)
    db.session.commit()
    rows = aggregate_rows(client_id=celine.id, warehouse_id=ny.id)
    upcs = {r["upc"] for r in rows}
    assert {available.upc, reserved.upc, packed.upc} <= upcs
    by_upc = {r["upc"]: r for r in rows}
    assert by_upc["UPC-AVAIL"]["on_hand"] == 1
    assert by_upc["UPC-RES"]["on_hand"] == 1
    assert by_upc["UPC-PACK"]["on_hand"] == 1


def test_d_e_shipped_only_hidden_unless_show_zero(app, db, admin_user, admin_client):
    celine, ny, _ = _world(db)
    live = _unit(celine, ny, "UPC-LIVE", "L-01")
    gone = _unit(celine, ny, "UPC-GONE", "Z-01")
    _ship(gone)
    db.session.commit()
    default_rows = aggregate_rows(client_id=celine.id, warehouse_id=ny.id)
    assert {r["upc"] for r in default_rows} == {"UPC-LIVE"}
    zero_rows = aggregate_rows(client_id=celine.id, warehouse_id=ny.id, include_zero=True)
    by_upc = {r["upc"]: r for r in zero_rows}
    assert by_upc["UPC-GONE"]["on_hand"] == 0
    assert by_upc["UPC-GONE"]["shipped"] == 1
    assert by_upc["UPC-LIVE"]["on_hand"] == 1

    html = admin_client.get(f"/inventory/?client_id={celine.id}&warehouse_id={ny.id}").get_data(as_text=True)
    assert "Show Zero On Hand" in html
    assert "UPC-LIVE" in html
    assert "UPC-GONE" not in html
    after = html.split('name="show_zero"', 1)[1][:80]
    assert "checked" not in after

    shown = admin_client.get(
        f"/inventory/?client_id={celine.id}&warehouse_id={ny.id}&show_zero=1"
    ).get_data(as_text=True)
    assert "UPC-GONE" in shown
    assert "UPC-LIVE" in shown
    assert live.upc and gone.upc


def test_f_g_kpis_ignore_depleted_even_with_show_zero(app, db, admin_user, admin_client):
    celine, ny, _ = _world(db)
    _unit(celine, ny, "UPC-LIVE", "L-01", sku="SKU-LIVE")
    gone = _unit(celine, ny, "UPC-GONE", "Z-99", sku="SKU-GONE")
    _ship(gone)
    db.session.commit()
    kpis = status_counts(client_id=celine.id, warehouse_id=ny.id)
    assert kpis["on_hand"] == 1
    assert kpis["unique_upcs"] == 1
    assert kpis["unique_skus"] == 1
    assert kpis["locations"] == 1
    html = admin_client.get(
        f"/inventory/?client_id={celine.id}&warehouse_id={ny.id}&show_zero=1"
    ).get_data(as_text=True)
    assert '<div class="num">1</div><div class="label">Unique UPCs</div>' in html
    assert '<div class="num">1</div><div class="label">Locations</div>' in html
    assert "UPC-GONE" in html


def test_h_search_follows_on_hand_rule(app, db, admin_user, admin_client):
    celine, ny, _ = _world(db)
    _unit(celine, ny, "UPC-LIVE", "L-01")
    gone = _unit(celine, ny, "UPC-GONE", "Z-01")
    _ship(gone)
    db.session.commit()
    live_hits = search_units("UPC-", client_id=celine.id, warehouse_id=ny.id)
    assert {u.upc for u in live_hits} == {"UPC-LIVE"}
    zero_hits = search_units("UPC-", client_id=celine.id, warehouse_id=ny.id, include_zero=True)
    assert {u.upc for u in zero_hits} == {"UPC-LIVE", "UPC-GONE"}

    html = admin_client.get(
        f"/inventory/search?client_id={celine.id}&warehouse_id={ny.id}&q=UPC-"
    ).get_data(as_text=True)
    assert "UPC-LIVE" in html
    assert "UPC-GONE" not in html
    shown = admin_client.get(
        f"/inventory/search?client_id={celine.id}&warehouse_id={ny.id}&q=UPC-&show_zero=1"
    ).get_data(as_text=True)
    assert "UPC-GONE" in shown


def test_i_j_upc_history_and_transactions_at_zero(app, db, admin_user, admin_client):
    celine, ny, _ = _world(db)
    gone = _unit(celine, ny, "UPC-GONE", "Z-01")
    _ship(gone)
    db.session.commit()
    resp = admin_client.get(
        f"/inventory/upc?client_id={celine.id}&warehouse_id={ny.id}&upc=UPC-GONE"
    )
    assert resp.status_code == 200
    html = resp.get_data(as_text=True)
    assert "UPC-GONE" in html
    assert "On Hand" in html
    assert ">0<" in html or "ON HAND" in html.upper()

    tx_html = admin_client.get(
        f"/inventory/transactions?client_id={celine.id}&warehouse_id={ny.id}&upc=UPC-GONE"
    ).get_data(as_text=True)
    assert "IMPORT" in tx_html
    assert "SHIP" in tx_html
    types = {
        t.transaction_type
        for t in InventoryTransaction.query.filter_by(upc="UPC-GONE")
    }
    assert {LedgerType.IMPORT, LedgerType.RESERVE, LedgerType.PACK, LedgerType.SHIP} <= types
    assert InventoryUnit.query.filter_by(upc="UPC-GONE").one().status == UnitStatus.SHIPPED


def test_k_l_excel_follows_show_zero(app, db, admin_user, admin_client):
    celine, ny, _ = _world(db)
    _unit(celine, ny, "UPC-LIVE", "L-01")
    gone = _unit(celine, ny, "UPC-GONE", "Z-01")
    _ship(gone)
    db.session.commit()

    default = admin_client.get(f"/inventory/export.xlsx?client_id={celine.id}&warehouse_id={ny.id}")
    assert default.status_code == 200
    book = load_workbook(filename=BytesIO(default.data))
    sheet = book.active
    headers = [cell.value for cell in sheet[1]]
    assert "ON_HAND" in headers
    upcs = [row[headers.index("UPC")] for row in sheet.iter_rows(min_row=2, values_only=True)]
    assert "UPC-LIVE" in upcs
    assert "UPC-GONE" not in upcs

    with_zero = admin_client.get(
        f"/inventory/export.xlsx?client_id={celine.id}&warehouse_id={ny.id}&show_zero=1"
    )
    book2 = load_workbook(filename=BytesIO(with_zero.data))
    sheet2 = book2.active
    headers2 = [cell.value for cell in sheet2[1]]
    upcs2 = [row[headers2.index("UPC")] for row in sheet2.iter_rows(min_row=2, values_only=True)]
    on_hands = {
        row[headers2.index("UPC")]: row[headers2.index("ON_HAND")]
        for row in sheet2.iter_rows(min_row=2, values_only=True)
    }
    assert "UPC-GONE" in upcs2
    assert on_hands["UPC-GONE"] == 0
    assert on_hands["UPC-LIVE"] == 1


def test_warehouse_zero_hidden_by_default(app, db, admin_user, admin_client):
    celine, ny, nj = _world(db)
    _unit(celine, ny, "UPC-NY", "A-01")
    gone = _unit(celine, nj, "UPC-NJ", "B-01")
    _ship(gone)
    db.session.commit()
    html = admin_client.get(f"/inventory/?client_id={celine.id}").get_data(as_text=True)
    assert "Warehouse Comparison" in html
    assert html.count("01-CEL-NY") >= 1
    # NJ has only shipped stock; hidden from comparison and default table
    assert "UPC-NJ" not in html
    shown = admin_client.get(f"/inventory/?client_id={celine.id}&show_zero=1").get_data(as_text=True)
    assert "UPC-NJ" in shown
