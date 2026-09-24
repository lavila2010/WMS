"""Order Processing Close Short: missing allocated units on the final carton."""

import threading

import pytest

from app.constants import (
    AllocationStatus,
    CartonStatus,
    InventoryIssueStatus,
    InventoryIssueType,
    LedgerType,
    OrderStatus,
    PickTicketStatus,
    UnitStatus,
)
from app.extensions import db
from app.models import (
    Allocation,
    AuditEvent,
    Carton,
    CartonContent,
    InventoryIssue,
    InventoryTransaction,
    InventoryUnit,
    Order,
    OrderLine,
    PickTicket,
    User,
)
from app.services.allocation import allocate_order, approve_partial_allocation
from app.services.documents import reconciliation_snapshot, render_closure_pdf
from app.services.end_of_day import build_eod_rows, export_eod_excel
from app.services.fulfillment import order_quantities
from app.services.inventory_issues import list_issues
from app.services.inventory_ledger import create_available_unit
from app.services.inventory_query import aggregate_rows, status_counts
from app.services.packing_list import render_packing_list_pdf
from app.services.pick_tickets import create_pick_ticket
from app.services.processing import (
    ProcessingError,
    can_close,
    can_close_short,
    close_order,
    close_order_short,
    remaining_rows,
    request_close,
    scan_upc,
    set_weight,
)
from tests.conftest import create_user, form_data, login
from tests.pdf_support import extract_pdf_text
from tests.test_v2_phase03_orders import _line
from tests.test_v2_phase04_allocation import _import_order, _world
from tests.test_v2_phase05_pick_tickets import _ready_order
from tests.test_v2_phase06_processing import _processing_order


def _pack_and_close_one(admin_user, order, carton, upc="UPC-A"):
    scan_upc(order, carton, upc, admin_user)
    request_close(carton, admin_user)
    set_weight(carton, 1.25, user=admin_user)
    return Order.query.get(order.id)


def _short_ready(db, admin_user, qty=2, packed=1, upc="UPC-A"):
    w, order, ticket, carton = _processing_order(db, admin_user, qty=qty, upc=upc)
    for _ in range(packed):
        scan_upc(order, carton, upc, admin_user)
    request_close(carton, admin_user)
    set_weight(carton, 2.0, user=admin_user)
    order = Order.query.get(order.id)
    ticket = PickTicket.query.get(ticket.id)
    return w, order, ticket, Carton.query.get(carton.id)


def test_normal_complete_processing_unchanged(app, db, admin_user):
    _, order, _, carton = _processing_order(db, admin_user, qty=2)
    scan_upc(order, carton, "UPC-A", admin_user)
    scan_upc(order, carton, "UPC-A", admin_user)
    request_close(carton, admin_user)
    set_weight(carton, 2.2, user=admin_user)
    close_order(Order.query.get(order.id), admin_user)
    order = Order.query.get(order.id)
    assert order.status == OrderStatus.CLOSED
    assert order.short_closed is False
    assert order.short_qty == 0
    assert sum(line.qty_shipped for line in order.lines) == 2
    assert sum(line.qty_short for line in order.lines) == 0
    assert InventoryUnit.query.filter_by(allocated_order_id=order.id, status=UnitStatus.SHIPPED).count() == 2
    assert InventoryIssue.query.filter_by(source_order_id=order.id).count() == 0


def test_reserved_remaining_without_confirmation_cannot_close(app, db, admin_user):
    _, order, _, carton = _processing_order(db, admin_user, qty=2)
    scan_upc(order, carton, "UPC-A", admin_user)
    request_close(carton, admin_user)
    set_weight(carton, 1.0, user=admin_user)
    order = Order.query.get(order.id)
    ok, reason = can_close(order)
    assert ok is False
    assert "Remaining reserved units must be zero" in reason
    with pytest.raises(ProcessingError, match="Remaining reserved units must be zero"):
        close_order(order, admin_user)
    with pytest.raises(ProcessingError, match="explicit confirmation"):
        close_order_short(order, admin_user, confirmed=False)
    assert InventoryUnit.query.filter_by(status=UnitStatus.RESERVED).count() == 1
    assert InventoryUnit.query.filter_by(status=UnitStatus.MISSING).count() == 0


def test_close_short_one_missing_reserved_to_missing(app, db, admin_user):
    _, order, ticket, _ = _short_ready(db, admin_user, qty=2, packed=1)
    document = close_order_short(order, admin_user, confirmed=True)
    order = Order.query.get(order.id)
    ticket = PickTicket.query.get(ticket.id)
    missing = InventoryUnit.query.filter_by(status=UnitStatus.MISSING).one()
    assert missing.status == UnitStatus.MISSING
    assert missing.allocated_order_id is None
    assert missing.allocation_id
    assert document is not None
    assert order.status == OrderStatus.CLOSED
    assert order.short_closed is True
    assert order.short_qty == 1
    assert ticket.status == PickTicketStatus.CLOSED
    assert ticket.short_closed is True
    assert ticket.short_qty == 1
    assert ticket.packed_qty == 1
    assert ticket.expected_qty == 2


def test_missing_unit_not_available_or_on_hand(app, db, admin_user):
    _, order, _, _ = _short_ready(db, admin_user, qty=2, packed=1)
    close_order_short(order, admin_user, confirmed=True)
    missing = InventoryUnit.query.filter_by(status=UnitStatus.MISSING).one()
    assert missing.status != UnitStatus.AVAILABLE
    assert missing.status not in UnitStatus.ON_HAND
    on_hand = InventoryUnit.query.filter(InventoryUnit.status.in_(UnitStatus.ON_HAND)).count()
    assert InventoryUnit.query.filter_by(status=UnitStatus.MISSING).count() == 1
    assert on_hand == InventoryUnit.query.filter(
        InventoryUnit.status.in_([UnitStatus.AVAILABLE, UnitStatus.RESERVED, UnitStatus.PACKED])
    ).count()
    counts = status_counts(client_id=missing.client_id, warehouse_id=missing.warehouse_id)
    assert counts["on_hand"] == counts["available"] + counts["reserved"] + counts["packed"]
    rows = aggregate_rows(client_id=missing.client_id, warehouse_id=missing.warehouse_id, include_zero=True)
    matching = [row for row in rows if row.get("upc") == missing.upc]
    for row in matching:
        assert row["on_hand"] == row["available"] + row["reserved"] + row["packed"]


def test_missing_unit_cannot_be_allocated_to_another_order(app, db, admin_user):
    w, order, _, _ = _short_ready(db, admin_user, qty=2, packed=1)
    close_order_short(order, admin_user, confirmed=True)
    missing = InventoryUnit.query.filter_by(status=UnitStatus.MISSING).one()
    other = _import_order(
        admin_user,
        w["celine"],
        w["cel_ecom"],
        [_line(order="8801", upc="UPC-A", qty=1)],
    )
    result = allocate_order(other, user=admin_user)
    assert result["reserved"] == 0
    db.session.refresh(missing)
    assert missing.status == UnitStatus.MISSING
    assert Allocation.query.filter_by(inventory_unit_id=missing.id, status=AllocationStatus.ACTIVE).count() == 0


def test_inventory_issue_created_and_visible_missing(app, db, admin_user):
    w, order, ticket, _ = _short_ready(db, admin_user, qty=2, packed=1)
    close_order_short(order, admin_user, confirmed=True)
    issue = InventoryIssue.query.filter_by(source_order_id=order.id).one()
    assert issue.issue_type == InventoryIssueType.PICK_UNIT_NOT_FOUND
    assert issue.status == InventoryIssueStatus.MISSING
    assert issue.source_pick_ticket_id == ticket.id
    assert issue.resolution_note == "Confirmed missing during Processing Close Short"
    visible = list_issues(
        client_id=w["celine"].id,
        warehouse_id=w["cel_ny"].id,
        status=InventoryIssueStatus.MISSING,
        upc="UPC-A",
        admin=True,
    )
    assert {row.id for row in visible} == {issue.id}
    by_q = list_issues(status=InventoryIssueStatus.MISSING, q=ticket.pick_ticket_number, admin=True)
    assert issue.id in {row.id for row in by_q}


def test_ledger_missing_confirmed_and_allocation_terminal(app, db, admin_user):
    _, order, _, _ = _short_ready(db, admin_user, qty=2, packed=1)
    close_order_short(order, admin_user, confirmed=True)
    missing = InventoryUnit.query.filter_by(status=UnitStatus.MISSING).one()
    types = [
        txn.transaction_type
        for txn in InventoryTransaction.query.filter_by(inventory_unit_id=missing.id).order_by(
            InventoryTransaction.id
        )
    ]
    assert LedgerType.ISSUE_HOLD in types
    assert LedgerType.MISSING_CONFIRMED in types
    alloc = Allocation.query.filter_by(inventory_unit_id=missing.id).one()
    assert alloc.status == AllocationStatus.MISSING
    assert alloc.status != AllocationStatus.ACTIVE


def test_qty_shipped_excludes_missing_and_qty_short_increments(app, db, admin_user):
    _, order, _, _ = _short_ready(db, admin_user, qty=2, packed=1)
    close_order_short(order, admin_user, confirmed=True)
    line = OrderLine.query.filter_by(order_id=order.id).one()
    assert line.qty_ordered == 2
    assert line.qty_shipped == 1
    assert line.qty_packed == 1
    assert line.qty_short == 1
    qty = order_quantities(Order.query.get(order.id))
    assert qty["shipped"] == 1
    assert qty["short"] == 1
    assert qty["remaining"] == 0


def test_packed_carton_qty_excludes_missing(app, db, admin_user):
    _, order, _, carton = _short_ready(db, admin_user, qty=2, packed=1)
    close_order_short(order, admin_user, confirmed=True)
    assert CartonContent.query.filter_by(carton_id=carton.id).count() == 1
    missing = InventoryUnit.query.filter_by(status=UnitStatus.MISSING).one()
    assert CartonContent.query.filter_by(inventory_unit_id=missing.id).count() == 0
    carton = Carton.query.get(carton.id)
    assert carton.status == CartonStatus.CLOSED
    assert carton.weight == 2.0


def test_order_closure_and_packing_list_short_docs(app, db, admin_user):
    _, order, ticket, carton = _short_ready(db, admin_user, qty=2, packed=1)
    close_order_short(order, admin_user, confirmed=True)
    order = Order.query.get(order.id)
    snap = reconciliation_snapshot(order)
    assert snap["ordered"] == 2
    assert snap["shipped"] == 1
    assert snap["short"] == 1
    assert snap["short_closed"] is True
    assert snap["result"] == "CLOSED SHORT"
    closure = render_closure_pdf(order)
    closure_text = extract_pdf_text(closure)
    assert "CLOSED SHORT" in closure_text
    assert "not fully fulfilled" in closure_text.lower()
    assert "ORDERED" in closure_text.upper()
    assert "SHIPPED" in closure_text.upper()
    assert "SHORT" in closure_text.upper()
    assert "MISSING" in closure_text.upper()
    packing = render_packing_list_pdf(order, ticket=PickTicket.query.get(ticket.id))
    packing_text = extract_pdf_text(packing)
    assert "FULFILLMENT EXCEPTIONS" in packing_text
    assert "Order closed short with confirmed missing inventory." in packing_text
    assert carton.carton_number in packing_text
    missing = InventoryUnit.query.filter_by(status=UnitStatus.MISSING).one()
    assert packing_text.lower().count("qty in carton") >= 1
    assert CartonContent.query.filter_by(inventory_unit_id=missing.id).count() == 0


def test_end_of_day_shows_short_quantity(app, db, admin_user):
    _, order, _, _ = _short_ready(db, admin_user, qty=2, packed=1)
    close_order_short(order, admin_user, confirmed=True)
    order = Order.query.get(order.id)
    rows = build_eod_rows([order])
    assert rows[0]["short"] == 1
    assert rows[0]["shipped"] == 1
    assert rows[0]["ordered"] == 2
    assert rows[0]["short_closed"] is True
    assert rows[0]["close_type"] == "CLOSED SHORT"
    data = export_eod_excel(rows)
    from io import BytesIO

    from openpyxl import load_workbook

    book = load_workbook(filename=BytesIO(data))
    assert book.sheetnames == ["End of Day"]
    values = None
    for row in book["End of Day"].iter_rows(values_only=True):
        if row and row[0] == "Client":
            continue
        if row and row[3] == order.wms_order_id:
            values = row
            break
    assert values is not None
    assert values[11] == 2
    assert values[12] == 1
    assert values[13] == 1
    assert values[9] == "CLOSED SHORT"


def test_concurrent_short_close_and_scan_cannot_double_process(app, db, admin_user):
    _, order, _, carton = _short_ready(db, admin_user, qty=2, packed=1)
    open_carton = Carton.query.filter_by(order_id=order.id).filter(Carton.status != CartonStatus.CLOSED).first()
    oid = order.id
    uid = admin_user.id
    cid = open_carton.id if open_carton else carton.id
    outcomes = []

    def closer():
        with app.app_context():
            try:
                close_order_short(db.session.get(Order, oid), db.session.get(User, uid), confirmed=True)
                outcomes.append("closed")
            except Exception as exc:
                outcomes.append(f"close:{exc}")

    def scanner():
        with app.app_context():
            try:
                target = db.session.get(Order, oid)
                box = db.session.get(Carton, cid)
                scan_upc(target, box, "UPC-A", db.session.get(User, uid))
                outcomes.append("scanned")
            except Exception as exc:
                outcomes.append(f"scan:{exc}")

    threads = [threading.Thread(target=closer), threading.Thread(target=scanner)]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()
    db.session.expire_all()
    unit_states = {unit.status for unit in InventoryUnit.query.all() if unit.upc == "UPC-A"}
    assert not ({UnitStatus.PACKED, UnitStatus.MISSING} <= unit_states and InventoryUnit.query.filter_by(status=UnitStatus.RESERVED).count() > 0)
    missing_count = InventoryUnit.query.filter_by(status=UnitStatus.MISSING).count()
    packed_or_shipped = InventoryUnit.query.filter(InventoryUnit.status.in_([UnitStatus.PACKED, UnitStatus.SHIPPED])).count()
    assert missing_count + packed_or_shipped >= 2
    assert InventoryTransaction.query.filter_by(transaction_type=LedgerType.MISSING_CONFIRMED).count() <= 1
    assert any(item.startswith("close:") or item == "closed" for item in outcomes)
    assert any(item.startswith("scan:") or item == "scanned" for item in outcomes)


def test_rollback_leaves_original_state_on_injected_failure(app, db, admin_user):
    _, order, _, _ = _short_ready(db, admin_user, qty=2, packed=1)
    reserved_id = InventoryUnit.query.filter_by(status=UnitStatus.RESERVED).one().id
    before_issues = InventoryIssue.query.count()
    before_short = OrderLine.query.filter_by(order_id=order.id).one().qty_short
    with pytest.raises(ProcessingError, match="injected failure"):
        close_order_short(order, admin_user, confirmed=True, _fail_after=1)
    db.session.expire_all()
    unit = db.session.get(InventoryUnit, reserved_id)
    assert unit.status == UnitStatus.RESERVED
    assert InventoryIssue.query.count() == before_issues
    assert InventoryTransaction.query.filter_by(transaction_type=LedgerType.MISSING_CONFIRMED).count() == 0
    assert OrderLine.query.filter_by(order_id=order.id).one().qty_short == before_short
    assert Order.query.get(order.id).status == OrderStatus.PROCESSING
    assert PickTicket.query.filter_by(order_id=order.id).one().status == PickTicketStatus.OPEN


def test_tenant_isolation_for_short_close_issues(app, db, admin_user, client):
    w, order, _, _ = _short_ready(db, admin_user, qty=2, packed=1)
    close_order_short(order, admin_user, confirmed=True)
    issue = InventoryIssue.query.filter_by(source_order_id=order.id).one()
    create_user("dio", perms=["INVENTORY_ISSUES_VIEW", "DASHBOARD_VIEW", "INVENTORY_VIEW"], clients=[w["dior"].id])
    login(client, "dio")
    html = client.get("/inventory/issues?status=MISSING").get_data(as_text=True)
    assert order.wms_order_id not in html
    assert client.get(f"/inventory/issues/{issue.id}").status_code == 404


def test_order_close_short_permission_enforced(app, db, admin_user):
    w, order, _, _ = _short_ready(db, admin_user, qty=2, packed=1)
    limited = create_user(
        "processor",
        perms=["PROCESSING_VIEW", "PROCESSING_EXECUTE", "BOX_CLOSE", "ORDER_CLOSE", "DASHBOARD_VIEW"],
        clients=[w["celine"].id],
    )
    limited_client = app.test_client()
    login(limited_client, "processor")
    page = limited_client.get(f"/processing/{order.id}/short-close")
    assert page.status_code == 200
    denied = limited_client.post(
        f"/processing/{order.id}/short-close",
        data=form_data(limited_client, {"confirm_missing": "1"}, path=f"/processing/{order.id}/short-close"),
    )
    assert denied.status_code == 403
    assert InventoryUnit.query.filter_by(status=UnitStatus.MISSING).count() == 0
    assert limited.has_permission("ORDER_CLOSE_SHORT") is False


def test_multi_wave_ticket_short_preserves_remaining(app, db, admin_user):
    w = _world(db)
    for _ in range(2):
        create_available_unit(client_id=w["celine"].id, warehouse_id=w["cel_ny"].id, upc="UPC-A", location="A-01", sku="SKU-A")
    db.session.commit()
    order = _import_order(admin_user, w["celine"], w["cel_ecom"], [_line(order="7701", upc="UPC-A", qty=3)])
    allocate_order(order, user=admin_user)
    order = Order.query.get(order.id)
    approve_partial_allocation(order, user=admin_user)
    ticket = create_pick_ticket(Order.query.get(order.id))
    from app.services.processing import acquire_lock, ensure_open_carton, set_dimensions

    order = Order.query.get(order.id)
    acquire_lock(order, admin_user)
    carton = ensure_open_carton(order, admin_user)
    set_dimensions(carton, 10, 8, 6)
    db.session.commit()
    scan_upc(order, carton, "UPC-A", admin_user)
    request_close(carton, admin_user)
    set_weight(carton, 1.1, user=admin_user)
    close_order_short(Order.query.get(order.id), admin_user, confirmed=True)
    order = Order.query.get(order.id)
    qty = order_quantities(order)
    assert order.status == OrderStatus.PARTIALLY_FULFILLED
    assert order.short_closed is False
    assert qty["shipped"] == 1
    assert qty["short"] == 1
    assert qty["remaining"] == 1
    assert PickTicket.query.get(ticket.id).status == PickTicketStatus.CLOSED
    create_available_unit(client_id=w["celine"].id, warehouse_id=w["cel_ny"].id, upc="UPC-A", location="A-02", sku="SKU-A")
    db.session.commit()
    allocate_order(Order.query.get(order.id), user=admin_user)
    order = Order.query.get(order.id)
    assert order_quantities(order)["remaining"] == 0
    assert order_quantities(order)["currently_allocated"] == 1


def test_processing_ui_shortage_card_and_confirm_checkbox(app, db, admin_user, admin_client):
    _, order, _, _ = _short_ready(db, admin_user, qty=2, packed=1)
    detail = admin_client.get(f"/processing/{order.id}").get_data(as_text=True)
    assert "have not been scanned" in detail
    assert "Remaining Units Not Scanned" in detail
    assert "Review Missing Units" in detail
    review = admin_client.get(f"/processing/{order.id}/short-close")
    html = review.get_data(as_text=True)
    assert review.status_code == 200
    assert "ORDER HAS MISSING UNITS" in html
    assert "confirm_missing" in html
    assert "Confirm &amp; Close Short" in html
    assert AuditEvent.query.filter_by(event_type="PROCESSING_SHORTAGE_DETECTED").count() >= 1
    no_check = admin_client.post(
        f"/processing/{order.id}/short-close",
        data=form_data(admin_client, {}, path=f"/processing/{order.id}/short-close"),
        follow_redirects=True,
    )
    assert InventoryUnit.query.filter_by(status=UnitStatus.MISSING).count() == 0
    assert "confirmation checkbox" in no_check.get_data(as_text=True)
    ok = admin_client.post(
        f"/processing/{order.id}/short-close",
        data=form_data(admin_client, {"confirm_missing": "1"}, path=f"/processing/{order.id}/short-close"),
        follow_redirects=True,
    )
    assert ok.status_code == 200
    assert InventoryUnit.query.filter_by(status=UnitStatus.MISSING).count() == 1
    assert AuditEvent.query.filter_by(event_type="PROCESSING_SHORT_CLOSE_CONFIRMED").count() >= 1
    assert AuditEvent.query.filter_by(event_type="ORDER_CLOSED_SHORT").count() >= 1
    assert AuditEvent.query.filter_by(event_type="PICK_TICKET_CLOSED_SHORT").count() >= 1


def test_remaining_rows_include_style_color_size(app, db, admin_user):
    _, order, _, carton = _processing_order(db, admin_user, qty=2)
    for unit in InventoryUnit.query.filter_by(allocated_order_id=order.id, status=UnitStatus.RESERVED):
        unit.style = "BAG"
        unit.color = "BLK"
        unit.size = "OS"
    db.session.commit()
    rows = remaining_rows(order)
    assert rows[0]["style"] == "BAG"
    assert rows[0]["color"] == "BLK"
    assert rows[0]["size"] == "OS"
    assert rows[0]["qty"] == 2
    ok, _ = can_close_short(order, admin_user)
    assert ok is False
    scan_upc(order, carton, "UPC-A", admin_user)
    request_close(carton, admin_user)
    set_weight(carton, 1.0, user=admin_user)
    ok, _ = can_close_short(Order.query.get(order.id), admin_user)
    assert ok is True
